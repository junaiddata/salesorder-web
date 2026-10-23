"""Customer list: add customers once (or import the Excel), then pick the name in a quotation to fill their details."""
import os
import re

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from .models import Customer

MAX_EXCEL_BYTES = 5 * 1024 * 1024
MAX_EXCEL_ROWS = 20000
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'sample_files', 'customer_list_template.xlsx')
HEADERS = {
    'customer id': 'code', 'customer code': 'code', 'id': 'code',
    'customer name': 'name', 'name': 'name',
    'trn number': 'trn', 'trn': 'trn',
    'contact number': 'contact_number', 'contact': 'contact_number', 'phone': 'contact_number',
    'contact person': 'contact_person',
    'address': 'address',
    'customer email': 'email', 'email': 'email', 'customer mail': 'email',
}
FIELDS = ('code', 'name', 'trn', 'contact_number', 'contact_person', 'address', 'email')
LIMITS = {'code': 50, 'name': 300, 'trn': 50, 'contact_number': 200, 'contact_person': 200, 'email': 300}


def _cell(value):
    """Cell text. Numbers come back without a trailing .0 (a TRN or phone typed as a number stays digits)."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


@login_required
def customer_list(request):
    q = (request.GET.get('q') or '').strip()
    customers = Customer.objects.all()
    for t in [t for t in re.split(r'\s+', q) if t]:
        customers = customers.filter(Q(name__icontains=t) | Q(code__icontains=t) | Q(trn__icontains=t) | Q(email__icontains=t)
                                     | Q(contact_person__icontains=t) | Q(contact_number__icontains=t))
    page = Paginator(customers, 50).get_page(request.GET.get('page'))
    return render(request, 'quotation/customer_list.html', {'page': page, 'q': q, 'total': Customer.objects.count(), 'shown': page.paginator.count})


@login_required
def customer_form(request, pk=None):
    customer = get_object_or_404(Customer, pk=pk) if pk else None
    errors = []
    values = {k: getattr(customer, k) if customer else '' for k in FIELDS}
    if request.method == 'POST':
        values = {k: (request.POST.get(k) or '').strip() for k in FIELDS}
        if not values['code']:
            errors.append('Customer ID is required.')
        elif Customer.objects.filter(code_key=values['code'].upper()).exclude(pk=customer.pk if customer else None).exists():
            errors.append(f'Customer ID "{values["code"]}" already exists in the list.')
        if not values['name']:
            errors.append('Customer Name is required.')
        for k, n in LIMITS.items():
            if len(values[k]) > n:
                errors.append(f'{k.replace("_", " ").title()} is too long (max {n} characters).')
        if not errors:
            customer = customer or Customer()
            for k in FIELDS:
                setattr(customer, k, values[k])
            customer.save()
            messages.success(request, f'Customer {customer.code} saved.')
            return redirect('quotation:customer_list')
    return render(request, 'quotation/customer_form.html', {'customer': customer, 'values': values, 'errors': errors})


@login_required
@require_POST
def customer_delete(request, pk):
    customer = get_object_or_404(Customer, pk=pk)
    label = customer.code
    customer.delete()
    messages.success(request, f'Customer {label} deleted. Quotations already made are not changed.')
    return redirect('quotation:customer_list')


@login_required
@require_POST
def customer_import_excel(request):
    """Read the customer Excel: a new Customer ID is added, an existing one is updated (a blank cell never wipes saved data)."""
    f = request.FILES.get('excel')
    if not f:
        messages.error(request, 'Please choose the Excel file.')
        return redirect('quotation:customer_list')
    if f.size > MAX_EXCEL_BYTES or not f.name.lower().endswith('.xlsx'):
        messages.error(request, 'Please upload an .xlsx Excel file (max 5 MB).')
        return redirect('quotation:customer_list')
    from openpyxl import load_workbook
    try:
        sheet = load_workbook(f, read_only=True, data_only=True).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True, max_row=MAX_EXCEL_ROWS + 1))
    except Exception:
        messages.error(request, 'This file could not be read. Please use the downloaded template (.xlsx).')
        return redirect('quotation:customer_list')
    cols = {}
    for i, h in enumerate(rows[0] if rows else ()):
        field = HEADERS.get(' '.join(_cell(h).lower().split()))
        if field and field not in cols:
            cols[field] = i
    if 'code' not in cols or 'name' not in cols:
        messages.error(request, 'The first row must have "Customer ID" and "Customer Name" columns. Please use the downloaded template.')
        return redirect('quotation:customer_list')

    existing = {c.code_key: c for c in Customer.objects.all()}
    created = updated = 0
    problems = []
    seen = set()
    for n, row in enumerate(rows[1:], start=2):
        data = {k: _cell(row[i]) if i < len(row) else '' for k, i in cols.items()}
        code, name = data.get('code', ''), data.get('name', '')
        if not code:
            if any(data.values()):
                problems.append(f'Row {n}: Customer ID is empty, skipped.')
            continue
        too_long = [k for k, lim in LIMITS.items() if len(data.get(k, '')) > lim]
        if too_long:
            problems.append(f'Row {n}: {code} has a too long {too_long[0].replace("_", " ")}, skipped.')
            continue
        key = code.upper()
        if key in seen:
            problems.append(f'Row {n}: {code} is repeated in the sheet, skipped.')
            continue
        seen.add(key)
        customer = existing.get(key)
        if not customer and not name:
            problems.append(f'Row {n}: {code} has no Customer Name, skipped.')
            continue
        is_new = customer is None
        customer = customer or Customer(code=code)
        for k in FIELDS[1:]:
            if data.get(k):
                setattr(customer, k, data[k])
        customer.save()
        existing[key] = customer
        created += is_new
        updated += not is_new
    messages.success(request, f'Excel imported: {created} new, {updated} updated, {len(problems)} skipped.')
    if problems:   # one message, not one per row
        more = f' …and {len(problems) - 5} more skipped rows.' if len(problems) > 5 else ''
        messages.warning(request, 'Skipped: ' + ' '.join(problems[:5]) + more)
    return redirect('quotation:customer_list')


@login_required
def customer_template(request):
    if not os.path.exists(TEMPLATE_PATH):
        raise Http404
    return FileResponse(open(TEMPLATE_PATH, 'rb'), as_attachment=True, filename='customer_list_template.xlsx')
