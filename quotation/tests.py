import os
import shutil
import tempfile
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from PIL import Image
from PyPDF2 import PdfReader

from .models import Customer, Product, product_key, Quotation, QuotationEmail, QuotationItem

TMP_ROOT = tempfile.mkdtemp(prefix='qtn_test_')


def make_pdf(text='Datasheet', pages=1):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for n in range(pages):
        c.drawString(100, 700, f'{text} page {n + 1}')
        c.showPage()
    c.save()
    return buf.getvalue()


PDF_BYTES = make_pdf()


def jpeg_size(data):
    assert data[:3] == b'\xff\xd8\xff', 'not a JPEG'
    return Image.open(BytesIO(data)).size


def png_file(name='p.png'):
    buf = BytesIO()
    Image.new('RGB', (40, 30), (200, 30, 30)).save(buf, 'PNG')
    return SimpleUploadedFile(name, buf.getvalue(), content_type='image/png')


@override_settings(QUOTATION_PRIVATE_ROOT=TMP_ROOT, QUOTATION_PUBLIC_BASE_URL='https://example.test',
                   QUOTATION_STAMP_USERS=())   # tests start from "everyone may stamp", whatever settings.py says
class QuotationTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP_ROOT, ignore_errors=True)

    def setUp(self):
        self.user = get_user_model().objects.create_user('tester', password='x', email='t@x.com')
        self.client.force_login(self.user)

    def _payload(self, **extra):
        data = {
            'date': '2026-10-20', 'to_text': 'ACME LLC\nDubai',
            'remarks': 'Deposit before production\n\n2) Port of loading: Shenzhen\n3. Net packing', 'validity': '30 days', 'stock_availability': '2026-11-05', 'project_reference': 'PRJ-77 Marina Tower',
            'country_of_origin': 'Italy', 'sales_person': 'Aiswarya', 'email': 'a@x.com',
            'contact_number': '0501234567',
            'row_key': ['n1', 'n2'], 'item_id': ['', ''],
            'item_no': ['H12', ''],
            'item_model': ['AR-50', 'PG-12'],
            'item_brand': ['ARISTON', 'PEGLER'], 'item_description': ['Heater 50L', 'Valve'],
            'item_size': ['50L', '1/2"'], 'item_color': ['White', ''],
            'item_qty': ['10', '2'], 'item_price': ['100.50', '25'],
            'item_warranty': ['2 Years', ''],
            'remove_photo': ['0', '0'], 'remove_ds': ['0', '0'],
        }
        data.update(extra)
        return data

    def test_create_with_files_and_totals(self):
        resp = self.client.post(reverse('quotation:new'), self._payload(
            photo_n1=png_file(), ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        self.assertRedirects(resp, reverse('quotation:list'))
        q = Quotation.objects.get()
        self.assertEqual(q.total, Decimal('1055.00'))
        self.assertEqual(q.vat_amount, Decimal('52.75'))        # 5% of 1,055.00
        self.assertEqual(q.grand_total, Decimal('1107.75'))
        first, second = q.items.all()
        self.assertTrue(first.photo and first.datasheet and first.share_token)
        self.assertFalse(second.datasheet or second.share_token)

    def test_public_datasheet_needs_no_login_and_only_serves_that_file(self):
        self.client.post(reverse('quotation:new'), self._payload(
            ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        token = QuotationItem.objects.get(brand='ARISTON').share_token
        self.client.logout()
        resp = self.client.get(reverse('quotation:datasheet', args=[token]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'image/jpeg')
        self.assertEqual(jpeg_size(b''.join(resp.streaming_content))[0], 1240)
        self.assertTrue(resp['Content-Disposition'].startswith('attachment'))   # the PDF link downloads the file
        viewed = self.client.get(reverse('quotation:datasheet', args=[token]) + '?view=1')
        self.assertTrue(viewed['Content-Disposition'].startswith('inline'))
        self.assertEqual(self.client.get(reverse('quotation:datasheet', args=['nope'])).status_code, 404)

    def test_other_pages_require_login(self):
        self.client.logout()
        for name, args in (('list', []), ('new', []), ('pdf', [1]), ('photo', [1])):
            resp = self.client.get(reverse(f'quotation:{name}', args=args))
            self.assertEqual(resp.status_code, 302, name)

    def test_rejects_non_pdf_datasheet_and_bad_image(self):
        resp = self.client.post(reverse('quotation:new'), self._payload(
            ds_n1=SimpleUploadedFile('sheet.pdf', b'not a pdf'),
            photo_n2=SimpleUploadedFile('x.png', b'garbage', 'image/png')))
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Quotation.objects.exists())
        self.assertContains(resp, 'PDF or an image')
        self.assertContains(resp, 'valid image')

    def test_no_items_is_rejected(self):
        resp = self.client.post(reverse('quotation:new'), self._payload(
            item_no=['', ''], item_model=['', ''], item_brand=['', ''], item_description=['', '']))
        self.assertContains(resp, 'at least one item')
        self.assertFalse(Quotation.objects.exists())

    def test_qtn_numbers_are_generated_per_year_in_order(self):
        for _ in range(3):
            self.client.post(reverse('quotation:new'), self._payload())
        self.client.post(reverse('quotation:new'), self._payload(date='2027-01-05'))   # a new year starts at 001
        self.client.post(reverse('quotation:new'), self._payload(date='2026-12-31'))
        self.assertEqual(list(Quotation.objects.order_by('id').values_list('qtn_no', flat=True)),
                         ['JG-2026-001', 'JG-2026-002', 'JG-2026-003', 'JG-2027-001', 'JG-2026-004'])

    def test_a_typed_qtn_no_is_ignored_and_old_numbers_do_not_disturb_the_sequence(self):
        Quotation.objects.create(qtn_no='20261016', date='2026-10-16')      # older, hand-typed numbers
        Quotation.objects.create(qtn_no='JG-2026-ABC', date='2026-10-16')
        Quotation.objects.create(qtn_no='JG-2026-009', date='2026-10-16')
        self.client.post(reverse('quotation:new'), self._payload(qtn_no='HACKED-1'))
        self.assertEqual(Quotation.objects.order_by('-id').first().qtn_no, 'JG-2026-010')
        self.assertFalse(Quotation.objects.filter(qtn_no='HACKED-1').exists())

    def test_number_grows_past_999_and_an_edit_keeps_its_number(self):
        Quotation.objects.create(qtn_no='JG-2026-999', date='2026-10-16')
        self.client.post(reverse('quotation:new'), self._payload())
        q = Quotation.objects.order_by('-id').first()
        self.assertEqual(q.qtn_no, 'JG-2026-1000')
        self.client.post(reverse('quotation:edit', args=[q.pk]), self._payload(
            qtn_no='JG-2026-001', date='2027-03-01', row_key=[f'e{i.pk}' for i in q.items.all()],
            item_id=[str(i.pk) for i in q.items.all()]))
        q.refresh_from_db()
        self.assertEqual((q.qtn_no, q.date.year), ('JG-2026-1000', 2027))   # the number never changes on edit

    def test_two_people_saving_at_once_get_different_numbers(self):
        from unittest import mock
        self.client.post(reverse('quotation:new'), self._payload())          # JG-2026-001 exists
        # Simulate the race: the first number tried was taken a moment ago by someone else.
        real = Quotation.next_qtn_no
        calls = []

        def stale_then_real(year):
            calls.append(year)
            return 'JG-2026-001' if len(calls) == 1 else real(year)

        with mock.patch.object(Quotation, 'next_qtn_no', side_effect=stale_then_real):
            resp = self.client.post(reverse('quotation:new'), self._payload())
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(sorted(Quotation.objects.values_list('qtn_no', flat=True)), ['JG-2026-001', 'JG-2026-002'])

    def test_form_shows_the_next_number_read_only(self):
        year = timezone.localdate().year
        self.assertContains(self.client.get(reverse('quotation:new')), f'value="JG-{year}-001"')   # nothing saved yet
        self.client.post(reverse('quotation:new'), self._payload(date=timezone.localdate().isoformat()))
        self.client.post(reverse('quotation:new'), self._payload())     # a 2026 one; only matters if today is in 2026
        page = self.client.get(reverse('quotation:new'))
        expected = 3 if year == 2026 else 2
        self.assertContains(page, f'value="JG-{year}-{expected:03d}"')
        self.assertNotContains(page, 'name="qtn_no"')
        q = Quotation.objects.order_by('id').first()
        self.assertContains(self.client.get(reverse('quotation:edit', args=[q.pk])), f'value="{q.qtn_no}"')

    def test_edit_keeps_files_and_token_and_replace_keeps_token(self):
        self.client.post(reverse('quotation:new'), self._payload(
            photo_n1=png_file(), ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        q = Quotation.objects.get()
        item = q.items.first()
        token, photo_name = item.share_token, item.photo.name
        # Edit without touching files: nothing is lost.
        edit = self._payload(row_key=[f'e{item.pk}', 'n1'], item_id=[str(item.pk), str(q.items.last().pk)])
        edit['row_key'] = [f'e{item.pk}', f'e{q.items.last().pk}']
        self.client.post(reverse('quotation:edit', args=[q.pk]), edit)
        item.refresh_from_db()
        self.assertEqual((item.share_token, item.photo.name), (token, photo_name))
        self.assertTrue(item.datasheet)
        # Replace the datasheet: same token, new file.
        new_pdf = make_pdf('Revised', pages=2)
        edit[f'ds_e{item.pk}'] = SimpleUploadedFile('new.pdf', new_pdf, 'application/pdf')
        self.client.post(reverse('quotation:edit', args=[q.pk]), edit)
        item.refresh_from_db()
        self.assertEqual(item.share_token, token)
        self.client.logout()
        resp = self.client.get(reverse('quotation:datasheet', args=[token]))
        w, h = jpeg_size(b''.join(resp.streaming_content))
        self.assertGreater(h, 1.5 * w)   # the 2-page PDF is now one tall JPEG

    def test_remove_datasheet_kills_link(self):
        self.client.post(reverse('quotation:new'), self._payload(
            ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        q = Quotation.objects.get()
        item, other = q.items.all()
        token = item.share_token
        edit = self._payload(row_key=[f'e{item.pk}', f'e{other.pk}'], item_id=[str(item.pk), str(other.pk)],
                             remove_ds=['1', '0'])
        self.client.post(reverse('quotation:edit', args=[q.pk]), edit)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('quotation:datasheet', args=[token])).status_code, 404)

    def test_delete_removes_files(self):
        self.client.post(reverse('quotation:new'), self._payload(
            photo_n1=png_file(), ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        item = QuotationItem.objects.get(brand='ARISTON')
        paths = [item.photo.path, item.datasheet.path]
        self.assertTrue(all(os.path.exists(p) for p in paths))
        self.client.post(reverse('quotation:delete', args=[item.quotation_id]))
        self.assertFalse(any(os.path.exists(p) for p in paths))
        self.assertFalse(Quotation.objects.exists())

    def test_pdf_is_on_letterhead_with_content_and_download_link(self):
        self.client.post(reverse('quotation:new'), self._payload(
            photo_n1=png_file(), ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        q = Quotation.objects.get()
        resp = self.client.get(reverse('quotation:pdf', args=[q.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        reader = PdfReader(BytesIO(resp.content))
        self.assertEqual(len(reader.pages), 1)
        text = reader.pages[0].extract_text()
        for needle in ('QUOTATION', 'JG-2026-001', '20.10.2026', 'ACME LLC', 'ARISTON', 'Heater 50L',
                       'H12', 'MODEL', 'AR-50', 'TERMS AND CONDITIONS', 'TOTAL (AED)', '1,055.00', 'VAT 5% (AED)', '52.75', 'GRAND TOTAL (AED)', '1,107.75', '1.  Deposit before production', '2.  Port of loading: Shenzhen', '3.  Net packing', 'Validity', '30 days', 'Delivery availability', '05.11.2026', 'Quotation no.', 'Project', 'PRJ-77 Marina Tower', 'BANK ACCOUNT DETAILS',
                       'A/C NO: 012001111025', 'A/C NAME: JUNAID GALLERIA TRADING LLC', 'AE88 0380 0000 1200 1111 025',
                       'Salesperson', 'Aiswarya', 'a@x.com', '0501234567', 'Download', 'Datasheet'):
            self.assertIn(needle, text, needle)
        item = q.items.first()
        uris = [a.get_object()['/A']['/URI'] for a in reader.pages[0].get('/Annots', [])]
        self.assertIn(f'https://example.test/quote-builder/ds/{item.share_token}/', uris)

    def test_image_and_multipage_pdf_uploads_become_jpeg(self):
        buf = BytesIO()
        Image.new('RGBA', (800, 600), (0, 0, 0, 0)).save(buf, 'PNG')   # transparent PNG
        self.client.post(reverse('quotation:new'), self._payload(
            ds_n1=SimpleUploadedFile('scan.png', buf.getvalue(), 'image/png'),
            ds_n2=SimpleUploadedFile('multi.pdf', make_pdf(pages=3), 'application/pdf')))
        one, two = Quotation.objects.get().items.all()
        self.client.logout()
        sizes = []
        for item in (one, two):
            self.assertTrue(item.datasheet.name.endswith('.jpg'), item.datasheet.name)
            resp = self.client.get(reverse('quotation:datasheet', args=[item.share_token]))
            self.assertEqual(resp['Content-Type'], 'image/jpeg')
            self.assertIn('.jpg', resp['Content-Disposition'])
            sizes.append(jpeg_size(b''.join(resp.streaming_content)))
        self.assertEqual(sizes[0], (800, 600))
        self.assertGreater(sizes[1][1], 3 * 1500)   # 3 stacked A4 pages

    def test_pdf_with_too_many_pages_rejected(self):
        resp = self.client.post(reverse('quotation:new'), self._payload(
            ds_n1=SimpleUploadedFile('big.pdf', make_pdf(pages=11), 'application/pdf')))
        self.assertContains(resp, 'max 10')
        self.assertFalse(Quotation.objects.exists())

    def test_rows_are_never_split_across_pages(self):
        n = 40
        descs = [f'START{i:02d} ' + 'long description words ' * 25 + f'END{i:02d}' for i in range(n)]
        payload = self._payload(
            row_key=[f'n{i}' for i in range(n)], item_id=[''] * n, item_no=[''] * n, item_model=[''] * n,
            item_brand=['ARISTON'] * n, item_description=descs,
            item_size=[''] * n, item_color=[''] * n, item_qty=['1'] * n, item_price=['10'] * n,
            item_warranty=[''] * n, remove_photo=['0'] * n, remove_ds=['0'] * n)
        self.client.post(reverse('quotation:new'), payload)
        q = Quotation.objects.get()
        reader = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[q.pk])).content))
        self.assertGreater(len(reader.pages), 3)
        texts = [p.extract_text().replace('\n', ' ') for p in reader.pages]
        for i in range(n):
            page_of_start = [k for k, t in enumerate(texts) if f'START{i:02d}' in t]
            page_of_end = [k for k, t in enumerate(texts) if f'END{i:02d}' in t]
            self.assertEqual(page_of_start, page_of_end, f'item {i} was split across pages')
        # Total, VAT and Grand Total stay together on one page, with the last item above them.
        vat_pages = [k for k, t in enumerate(texts) if 'VAT 5% (AED)' in t]
        self.assertEqual(vat_pages, [k for k, t in enumerate(texts) if 'GRAND TOTAL (AED)' in t])
        self.assertEqual(len(vat_pages), 1)
        self.assertIn(f'START{n - 1:02d}', texts[vat_pages[0]])
        # the words (40 items x 10 = 400 + 5% VAT = 420) stay with the totals, and carry no label or currency name
        self.assertEqual([k for k, t in enumerate(texts) if 'AED : FOUR HUNDRED AND TWENTY ONLY' in t], vat_pages)
        self.assertFalse(any('AMOUNT IN WORDS' in t or 'DIRHAMS' in t for t in texts))

    def test_pdf_survives_many_items_and_a_description_longer_than_a_page(self):
        long_desc = 'FIRSTWORD ' + 'Concealed flushing cistern with steel support frame and fixing set. ' * 60 + 'LASTWORD'   # ~4,000 chars
        payload = self._payload(
            row_key=[f'n{i}' for i in range(30)], item_id=[''] * 30, item_no=[''] * 30, item_model=[''] * 30,
            item_brand=['ARISTON'] * 30, item_description=[long_desc if i == 5 else f'Item {i}' for i in range(30)],
            item_size=[''] * 30, item_color=[''] * 30, item_qty=['1'] * 30, item_price=['10'] * 30,
            item_warranty=[''] * 30, remove_photo=['0'] * 30, remove_ds=['0'] * 30)
        self.client.post(reverse('quotation:new'), payload)
        q = Quotation.objects.get()
        resp = self.client.get(reverse('quotation:pdf', args=[q.pk]))
        self.assertEqual(resp.status_code, 200)
        reader = PdfReader(BytesIO(resp.content))
        self.assertGreater(len(reader.pages), 2)
        everything = '\n'.join(p.extract_text() for p in reader.pages)
        for needle in ('TOTAL (AED)', '300.00', 'FIRSTWORD', 'LASTWORD'):   # nothing of the long text is lost
            self.assertIn(needle, everything)
        flat = ' '.join(everything.split())   # the repeated table header can fall inside a phrase, so count words
        for word in ('Concealed', 'cistern', 'support', 'fixing', 'set.'):
            self.assertEqual(flat.count(word), 60, word)
        last = reader.pages[-1].extract_text()
        for needle in ('BANK ACCOUNT DETAILS', 'IBAN'):
            self.assertIn(needle, last)

    def test_stock_availability_is_a_date_and_project_reference_is_optional(self):
        resp = self.client.post(reverse('quotation:new'), self._payload(stock_availability='soon'))
        self.assertContains(resp, 'valid Stock Availability date')
        self.assertFalse(Quotation.objects.exists())
        # Both optional: blank date and blank project reference save fine and print nothing for them.
        self.client.post(reverse('quotation:new'), self._payload(stock_availability='', project_reference=''))
        q = Quotation.objects.get()
        self.assertIsNone(q.stock_availability)
        reader = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[q.pk])).content))
        self.assertNotIn('PRJ-77', reader.pages[0].extract_text())
        # The edit form shows the saved date in the date box.
        self.client.post(reverse('quotation:edit', args=[q.pk]), self._payload(
            stock_availability='2026-12-01', row_key=[f'e{i.pk}' for i in q.items.all()],
            item_id=[str(i.pk) for i in q.items.all()]))
        page = self.client.get(reverse('quotation:edit', args=[q.pk]))
        self.assertContains(page, 'type="date" name="stock_availability"')
        self.assertContains(page, 'value="2026-12-01"')

    def test_pdf_orientation_option(self):
        self.client.post(reverse('quotation:new'), self._payload(
            ds_n1=SimpleUploadedFile('sheet.pdf', PDF_BYTES, 'application/pdf')))
        q = Quotation.objects.get()
        url = reverse('quotation:pdf', args=[q.pk])
        item = q.items.first()
        for query, landscape in (('', False), ('?orientation=portrait', False), ('?orientation=landscape', True),
                                 ('?orientation=junk', False)):
            resp = self.client.get(url + query)
            self.assertEqual(resp.status_code, 200, query)
            pdf_pages = PdfReader(BytesIO(resp.content)).pages
            page = pdf_pages[0]
            w, h = float(page.mediabox.width), float(page.mediabox.height)
            self.assertEqual(w > h, landscape, query)
            self.assertAlmostEqual(max(w, h), 841.9, delta=0.5)   # always A4
            self.assertAlmostEqual(min(w, h), 595.3, delta=0.5)
            text = '\n'.join(p.extract_text() for p in pdf_pages)   # the closing blocks may move to a 2nd page
            for needle in ('QUOTATION', 'JG-2026-001', 'H12', 'TERMS AND CONDITIONS', 'IBAN', 'Download'):
                self.assertIn(needle, text, (query, needle))
            uris = [a.get_object()['/A']['/URI'] for a in page.get('/Annots', [])]
            self.assertIn(f'https://example.test/quote-builder/ds/{item.share_token}/', uris, query)   # links survive
            self.assertIn('Landscape' if landscape else 'Quotation_JG-2026-001.pdf', resp['Content-Disposition'])

    def test_landscape_rows_still_never_split_and_survive_huge_descriptions(self):
        n = 30
        descs = [f'START{i:02d} ' + 'long description words ' * 25 + f'END{i:02d}' for i in range(n)]
        descs[7] = 'FIRSTWORD ' + 'Concealed flushing cistern with steel support frame and fixing set. ' * 60 + 'LASTWORD'
        payload = self._payload(
            row_key=[f'n{i}' for i in range(n)], item_id=[''] * n, item_no=[''] * n, item_model=[''] * n,
            item_brand=['ARISTON'] * n, item_description=descs,
            item_size=[''] * n, item_color=[''] * n, item_qty=['1'] * n, item_price=['10'] * n,
            item_warranty=[''] * n, remove_photo=['0'] * n, remove_ds=['0'] * n)
        self.client.post(reverse('quotation:new'), payload)
        q = Quotation.objects.get()
        resp = self.client.get(reverse('quotation:pdf', args=[q.pk]) + '?orientation=landscape')
        self.assertEqual(resp.status_code, 200)
        reader = PdfReader(BytesIO(resp.content))
        texts = [p.extract_text().replace('\n', ' ') for p in reader.pages]
        for i in range(n):
            if i == 7:
                continue
            self.assertEqual([k for k, t in enumerate(texts) if f'START{i:02d}' in t],
                             [k for k, t in enumerate(texts) if f'END{i:02d}' in t], f'item {i} was split')
        everything = ' '.join(' '.join(texts).split())
        self.assertIn('FIRSTWORD', everything)
        self.assertIn('LASTWORD', everything)
        self.assertIn('IBAN', texts[-1])

    def test_vat_is_five_percent_rounded_to_the_cent_and_amounts_add_up(self):
        q = Quotation.objects.create(qtn_no='V-1', date='2026-10-20')
        QuotationItem.objects.create(quotation=q, position=0, brand='A', qty=Decimal('0.50'), unit_price=Decimal('0.01'))   # 0.005
        QuotationItem.objects.create(quotation=q, position=1, brand='B', qty=Decimal('10'), unit_price=Decimal('0.01'))      # 0.10
        a, b = q.items.all()
        self.assertEqual((a.amount, b.amount), (Decimal('0.01'), Decimal('0.10')))   # a line rounds half up to the cent
        self.assertEqual(q.total, Decimal('0.11'))                                    # = the sum of the printed lines
        self.assertEqual(q.vat_amount, Decimal('0.01'))                               # 0.0055 -> 0.01
        self.assertEqual(q.grand_total, Decimal('0.12'))
        a.delete()
        self.assertEqual(q.vat_amount, Decimal('0.01'))                               # 0.005 rounds half UP, not to even
        empty = Quotation.objects.create(qtn_no='V-2', date='2026-10-20')
        self.assertEqual((empty.total, empty.vat_amount, empty.grand_total), (Decimal('0'), Decimal('0.00'), Decimal('0.00')))

    def test_totals_show_in_list_and_form(self):
        self.client.post(reverse('quotation:new'), self._payload())
        q = Quotation.objects.get()
        listing = self.client.get(reverse('quotation:list'))
        self.assertContains(listing, 'Grand Total (AED)')
        self.assertContains(listing, '1107.75')
        form = self.client.get(reverse('quotation:edit', args=[q.pk]))
        for needle in ('id="tTotal"', 'id="tVat"', 'id="tGrand"', 'VAT 5%'):
            self.assertContains(form, needle)

    # ---- company stamp -------------------------------------------------------------------------------------
    def _upload_stamp(self, name='stamp.png', **kw):
        buf = BytesIO()
        Image.new('RGB', (200, 200), (240, 240, 240)).save(buf, 'PNG')
        return self.client.post(reverse('quotation:stamp_settings'),
                                {'stamp': SimpleUploadedFile(name, buf.getvalue(), 'image/png'), **kw})

    @staticmethod
    def _pdf_image_count(resp):
        import fitz
        return sum(len(page.get_images()) for page in fitz.open(stream=resp.content, filetype='pdf'))   # all pages

    def test_stamp_upload_replace_remove_and_validation(self):
        from .models import QuotationStamp
        self.assertContains(self.client.get(reverse('quotation:stamp_settings')), 'No stamp uploaded')
        self._upload_stamp()
        stamp = QuotationStamp.get_instance()
        first = stamp.stamp.path
        self.assertTrue(os.path.exists(first))
        self.assertContains(self.client.get(reverse('quotation:stamp_settings')), 'Stamp uploaded')
        self.assertNotContains(self.client.get(reverse('quotation:stamp_settings')), first)   # never shown or linked
        self._upload_stamp('second.png')                                                       # replaces the old file
        stamp.refresh_from_db()
        self.assertFalse(os.path.exists(first))
        self.assertTrue(os.path.exists(stamp.stamp.path))
        bad = self.client.post(reverse('quotation:stamp_settings'),
                               {'stamp': SimpleUploadedFile('x.png', b'not an image', 'image/png')}, follow=True)
        self.assertContains(bad, 'valid image')
        stamp.refresh_from_db()
        self.assertTrue(stamp.stamp)                                                          # a bad upload keeps the old one
        self.client.post(reverse('quotation:stamp_settings'), {'remove_stamp': '1'})
        stamp.refresh_from_db()
        self.assertFalse(stamp.stamp)

    def test_stamp_page_requires_login_and_the_optional_user_restriction(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse('quotation:stamp_settings')).status_code, 302)
        self.client.force_login(self.user)
        with override_settings(QUOTATION_STAMP_USERS=['md', 'manager']):
            self.assertEqual(self.client.get(reverse('quotation:stamp_settings')).status_code, 403)
            self.assertEqual(self._upload_stamp().status_code, 403)
        with override_settings(QUOTATION_STAMP_USERS=['Tester']):                      # case-insensitive
            self.assertEqual(self.client.get(reverse('quotation:stamp_settings')).status_code, 200)

    def test_add_stamp_prints_the_stamp_inside_the_bank_box_only_when_ticked(self):
        self._upload_stamp()
        self.client.post(reverse('quotation:new'), self._payload())                           # not ticked
        plain = Quotation.objects.get()
        self.assertFalse(plain.add_stamp)
        url = reverse('quotation:pdf', args=[plain.pk])
        self.client.post(reverse('quotation:new'), self._payload(add_stamp='on'))
        stamped = Quotation.objects.order_by('-id').first()
        self.assertTrue(stamped.add_stamp)
        for query in ('', '?orientation=landscape'):
            without = self._pdf_image_count(self.client.get(url + query))
            resp = self.client.get(reverse('quotation:pdf', args=[stamped.pk]) + query)
            self.assertEqual(resp.status_code, 200)
            self.assertEqual(self._pdf_image_count(resp), without + 1, query)                 # exactly one extra image
            self.assertIn('BANK ACCOUNT DETAILS', ' '.join(p.extract_text() for p in PdfReader(BytesIO(resp.content)).pages))

    def test_add_stamp_without_an_uploaded_stamp_still_makes_a_pdf(self):
        self.client.post(reverse('quotation:new'), self._payload(add_stamp='on'))
        q = Quotation.objects.get()
        page = self.client.get(reverse('quotation:new'))
        self.assertContains(page, 'No stamp uploaded yet')
        resp = self.client.get(reverse('quotation:pdf', args=[q.pk]))
        self.assertEqual(resp.status_code, 200)

    def test_users_who_may_not_stamp_get_unstamped_pdfs_and_cannot_change_the_choice(self):
        self._upload_stamp()
        self.client.post(reverse('quotation:new'), self._payload(add_stamp='on'))
        q = Quotation.objects.get()
        base = self._pdf_image_count(self.client.get(reverse('quotation:pdf', args=[q.pk])))
        with override_settings(QUOTATION_STAMP_USERS=['md']):
            self.assertEqual(self._pdf_image_count(self.client.get(reverse('quotation:pdf', args=[q.pk]))), base - 1)
            edit = self._payload(row_key=[f'e{i.pk}' for i in q.items.all()], item_id=[str(i.pk) for i in q.items.all()])
            self.client.post(reverse('quotation:edit', args=[q.pk]), edit)                   # the tick box is absent
            q.refresh_from_db()
            self.assertTrue(q.add_stamp)                                                      # their edit left it as it was
            self.assertContains(self.client.get(reverse('quotation:edit', args=[q.pk])), 'only authorised users')
        # An authorised user can untick it.
        self.client.post(reverse('quotation:edit', args=[q.pk]), edit)
        q.refresh_from_db()
        self.assertFalse(q.add_stamp)

    def test_the_to_box_shows_a_label_for_every_detail_in_order(self):
        # nothing in the four customer fields (older style, only the free text): the labels are still there, and the
        # free text is printed under them
        self.client.post(reverse('quotation:new'), self._payload())
        old = Quotation.objects.get()
        text = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[old.pk])).content)).pages[0].extract_text()
        flat = ' '.join(text.split())
        self.assertIn('TO: Customer name Contact person Contact number Address TRN number Customer ID Customer email ACME LLC Dubai', flat)
        # all four filled: TO -> customer name -> contact -> address -> TRN number -> other details
        self.client.post(reverse('quotation:new'), self._payload(
            customer_name='A&B <Trading> LLC', customer_contact='Mr Ahmed, 050 123 4567',
            customer_address='PO BOX 34862\nDubai Investment Park 2', customer_trn='100123456700003',
            to_text='Other line'))
        q = Quotation.objects.order_by('-id').first()
        self.assertEqual((q.customer_name, q.customer_trn), ('A&B <Trading> LLC', '100123456700003'))
        for query in ('', '?orientation=landscape'):
            text = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[q.pk]) + query).content)).pages[0].extract_text()
            flat = ' '.join(text.split())
            order = ['TO:', 'Customer name A&B <Trading> LLC', 'Contact number Mr Ahmed, 050 123 4567',
                     'Address PO BOX 34862 Dubai Investment Park 2', 'TRN number 100123456700003', 'Other line']
            positions = [flat.index(x) for x in order]
            self.assertEqual(positions, sorted(positions), query)
        # the edit form and the list show them
        edit = self.client.get(reverse('quotation:edit', args=[q.pk]))
        for needle in ('name="customer_name"', 'name="customer_contact"', 'name="customer_address"', 'name="customer_trn"',
                       'A&amp;B &lt;Trading&gt; LLC', '100123456700003'):
            self.assertContains(edit, needle)
        listing = self.client.get(reverse('quotation:list'))
        self.assertContains(listing, 'A&amp;B &lt;Trading&gt; LLC')    # the list shows the customer name
        self.assertContains(listing, 'ACME LLC')                        # ... or the older free text when there is no name

    def test_terms_box_has_validity_delivery_availability_country_of_origin_and_warranty_in_that_order(self):
        self.client.post(reverse('quotation:new'), self._payload(warranty='Two years manufacturer warranty'))
        q = Quotation.objects.get()
        self.assertEqual(q.warranty, 'Two years manufacturer warranty')
        for query in ('', '?orientation=landscape'):
            pdf = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[q.pk]) + query).content))
            text = '\n'.join(p.extract_text() for p in pdf.pages)
            order = ['TERMS AND CONDITIONS', '1.  Deposit before production', '3.  Net packing',
                     'VALIDITY\n30 days', 'DELIVERY AVAILABILITY\n05.11.2026', 'COUNTRY OF ORIGIN\nItaly',
                     'WARRANTY\nTwo years manufacturer warranty', 'BANK ACCOUNT DETAILS']
            positions = [text.index(x) for x in order]
            self.assertEqual(positions, sorted(positions), query)
            self.assertEqual(text.count('TERMS AND CONDITIONS'), 1, query)       # one box, not two
            self.assertNotIn('STOCK AVAILABILITY', text)                         # it is called Delivery availability now
            # ... and they are still in the top block as well (Title-case labels there, capitals in the terms box)
            self.assertIn('Validity', text)
            self.assertIn('Delivery availability', text)
        page = self.client.get(reverse('quotation:edit', args=[q.pk]))
        self.assertContains(page, 'name="warranty"')
        self.assertContains(page, 'Two years manufacturer warranty')

    def test_top_block_has_three_sections_and_the_bottom_contact_box_is_gone(self):
        self.client.post(reverse('quotation:new'), self._payload(
            customer_name='ACME TRADING LLC', customer_contact='Mr Ahmed', customer_address='Dubai', customer_trn='100123',
            customer_ref='RFQ-8841'))
        q = Quotation.objects.get()
        self.assertEqual(q.customer_ref, 'RFQ-8841')
        for query in ('', '?orientation=landscape'):
            text = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[q.pk]) + query).content)).pages[0].extract_text()
            flat = ' '.join(text.split())
            # section 1: To
            for x in ('TO:', 'Customer name ACME TRADING LLC', 'Contact number Mr Ahmed', 'Address Dubai', 'TRN number 100123'):
                self.assertIn(x, flat, (query, x))
            # section 2: reference details, with the salesperson / email / contact that used to be in the bottom box
            for x in ('Quotation no. JG-2026-001', 'Customer ref. RFQ-8841', 'Salesperson Aiswarya', 'Email a@x.com',
                      'Contact 0501234567', 'Project reference PRJ-77 Marina Tower'):
                self.assertIn(x, flat, (query, x))
            # section 3: date, validity, delivery availability
            for x in ('Date 20.10.2026', 'Validity 30 days', 'Delivery availability 05.11.2026'):
                self.assertIn(x, flat, (query, x))
            # all of that is above the items table; nothing of the old bottom CONTACT box is left
            self.assertLess(flat.index('Delivery availability'), flat.index('ITEM NO.'))
            self.assertLess(flat.index('Project reference'), flat.index('ITEM NO.'))
            self.assertNotIn('CONTACT NUMBER', flat)
            self.assertNotIn('SALES PERSON', flat)
            self.assertEqual(flat.count('Salesperson'), 1)
        page = self.client.get(reverse('quotation:edit', args=[q.pk]))
        self.assertContains(page, 'name="customer_ref"')
        self.assertContains(page, 'Delivery Availability')

    def test_top_block_copes_with_nothing_filled_in_old_quotations_and_very_long_values(self):
        for ori in ('', '?orientation=landscape'):
            # a bare quotation: only the mandatory fields
            bare = Quotation.objects.create(qtn_no=f'BARE{len(ori)}', date='2026-10-20')
            QuotationItem.objects.create(quotation=bare, position=0, brand='A', qty=Decimal('1'), unit_price=Decimal('10'))
            resp = self.client.get(reverse('quotation:pdf', args=[bare.pk]) + ori)
            self.assertEqual(resp.status_code, 200)
            text = ' '.join(PdfReader(BytesIO(resp.content)).pages[0].extract_text().split())
            for x in ('TO:', f'Quotation no. BARE{len(ori)}', 'Salesperson', 'Date 20.10.2026', 'Validity', 'Delivery availability'):
                self.assertIn(x, text, (ori, x))
            # an older quotation that only has the free-text To box
            old = Quotation.objects.create(qtn_no=f'OLD{len(ori)}', date='2026-10-20', to_text='OLD STYLE CO\nPO BOX 1')
            QuotationItem.objects.create(quotation=old, position=0, brand='A', qty=Decimal('1'), unit_price=Decimal('10'))
            text = ' '.join(PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[old.pk]) + ori).content)).pages[0].extract_text().split())
            self.assertIn('OLD STYLE CO PO BOX 1', text)
            # very long values: an unbroken email, a long reference and address wrap instead of overflowing
            big = Quotation.objects.create(
                qtn_no=f'BIG{len(ori)}', date='2026-10-20', customer_name='X' * 150, customer_address='Y' * 200 + '\n' + 'word ' * 60,
                email='very.long.email.address.of.the.salesperson@a-really-long-company-domain-name.example.com',
                project_reference='PROJECT ' * 30, customer_ref='R' * 120, validity='V' * 90)
            QuotationItem.objects.create(quotation=big, position=0, brand='A', qty=Decimal('1'), unit_price=Decimal('10'))
            resp = self.client.get(reverse('quotation:pdf', args=[big.pk]) + ori)
            self.assertEqual(resp.status_code, 200)
            reader = PdfReader(BytesIO(resp.content))
            text = ''.join(reader.pages[0].extract_text().split())
            self.assertIn('very.long.email.address.of.the.salesperson@a-really-long-company-domain-name.example.com', text)
            self.assertIn('R' * 120, text)

    def test_blank_terms_facts_are_left_out_and_an_empty_box_is_not_printed(self):
        self.client.post(reverse('quotation:new'), self._payload(
            remarks='', validity='', stock_availability='', country_of_origin='', warranty='Only this'))
        q = Quotation.objects.get()
        text = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[q.pk])).content)).pages[0].extract_text()
        self.assertIn('TERMS AND CONDITIONS', text)
        self.assertIn('Only this', text)
        for label in ('COUNTRY OF ORIGIN', 'VALIDITY', 'DELIVERY AVAILABILITY'):
            self.assertNotIn(label, text)                    # blank ones are left out of the terms box
        self.client.post(reverse('quotation:new'), self._payload(
            remarks='', validity='', stock_availability='', country_of_origin='', warranty=''))
        empty = Quotation.objects.order_by('-id').first()
        text = PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[empty.pk])).content)).pages[0].extract_text()
        self.assertNotIn('TERMS AND CONDITIONS', text)
        self.assertIn('BANK ACCOUNT DETAILS', text)

    # ---- discount, total excl. VAT, VAT, grand total, amount in words ------------------------------------
    def test_amount_in_words(self):
        from .words import amount_in_words
        cases = {
            '0': 'ZERO ONLY',
            '0.50': 'ZERO AND FIFTY FILS ONLY',
            '21': 'TWENTY ONE ONLY',
            '100': 'ONE HUNDRED ONLY',
            '101': 'ONE HUNDRED AND ONE ONLY',
            '1005': 'ONE THOUSAND AND FIVE ONLY',
            '1050': 'ONE THOUSAND AND FIFTY ONLY',
            '1100': 'ONE THOUSAND ONE HUNDRED ONLY',
            '1107.75': 'ONE THOUSAND ONE HUNDRED AND SEVEN AND SEVENTY FIVE FILS ONLY',
            '123456.78': 'ONE HUNDRED AND TWENTY THREE THOUSAND FOUR HUNDRED AND FIFTY SIX AND SEVENTY EIGHT FILS ONLY',
            '2000000': 'TWO MILLION ONLY',
            '0.005': 'ZERO AND ONE FILS ONLY',      # half a fils rounds up to the cent
        }
        for value, words in cases.items():
            self.assertEqual(amount_in_words(Decimal(value)), words, value)

    def test_discount_amount_and_percent_calculations(self):
        q = Quotation.objects.create(qtn_no='D-1', date='2026-10-20')
        QuotationItem.objects.create(quotation=q, position=0, brand='A', qty=Decimal('10'), unit_price=Decimal('105.50'))   # 1,055.00
        self.assertEqual((q.discount_amount, q.net_total, q.vat_amount, q.grand_total),
                         (Decimal('0'), Decimal('1055.00'), Decimal('52.75'), Decimal('1107.75')))     # no discount: as before
        q.discount_type, q.discount_value = 'amount', Decimal('55')
        self.assertEqual((q.discount_amount, q.net_total, q.vat_amount, q.grand_total),
                         (Decimal('55'), Decimal('1000.00'), Decimal('50.00'), Decimal('1050.00')))    # VAT is on the discounted total
        q.discount_type, q.discount_value = 'percent', Decimal('10')
        self.assertEqual((q.discount_amount, q.net_total, q.vat_amount, q.grand_total),
                         (Decimal('105.50'), Decimal('949.50'), Decimal('47.48'), Decimal('996.98')))  # 47.475 rounds half up
        q.discount_type, q.discount_value = 'amount', Decimal('5000')                                   # more than the total: capped
        self.assertEqual((q.discount_amount, q.net_total, q.grand_total), (Decimal('1055.00'), Decimal('0.00'), Decimal('0.00')))
        q.discount_type, q.discount_value = 'percent', Decimal('250')
        self.assertEqual(q.discount_amount, Decimal('1055.00'))
        q.discount_value = Decimal('-5')
        self.assertEqual(q.discount_amount, Decimal('0'))

    def test_discount_validation_and_saving(self):
        url = reverse('quotation:new')
        for kwargs, message in (({'discount_type': 'amount', 'discount_value': '-5'}, 'cannot be negative'),
                                ({'discount_type': 'percent', 'discount_value': '101'}, 'more than 100%'),
                                ({'discount_type': 'amount', 'discount_value': '1055.01'}, 'more than the total')):
            resp = self.client.post(url, self._payload(**kwargs))
            self.assertContains(resp, message)
            self.assertFalse(Quotation.objects.exists())
        # exactly the total is allowed; blank means no discount; an unknown type falls back to AED
        self.client.post(url, self._payload(discount_type='amount', discount_value='1055'))
        self.client.post(url, self._payload(discount_type='percent', discount_value='12.5'))
        self.client.post(url, self._payload(discount_type='weird', discount_value=''))
        full, pct, none = Quotation.objects.order_by('id')
        self.assertEqual((full.discount_type, full.discount_value, full.grand_total), ('amount', Decimal('1055.00'), Decimal('0.00')))
        self.assertEqual((pct.discount_type, pct.discount_value), ('percent', Decimal('12.50')))
        self.assertEqual((none.discount_type, none.discount_value, none.grand_total), ('amount', Decimal('0'), Decimal('1107.75')))
        # the edit form shows the saved discount again
        page = self.client.get(reverse('quotation:edit', args=[pct.pk]))
        self.assertContains(page, '<option value="percent" selected>')
        self.assertContains(page, 'name="discount_value"')
        self.assertContains(page, 'value="12.50"')
        # 1,055.00 - 12.5% (131.875 -> 131.88) = 923.12; VAT 46.156 -> 46.16; grand 969.28 is what the list shows
        self.assertEqual(pct.grand_total, Decimal('969.28'))
        self.assertContains(self.client.get(reverse('quotation:list')), '969.28')

    def test_pdf_totals_block_with_and_without_discount(self):
        # without a discount: total excl. VAT, VAT, grand total, words
        self.client.post(reverse('quotation:new'), self._payload())
        plain = Quotation.objects.get()
        for query in ('', '?orientation=landscape'):
            text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(
                self.client.get(reverse('quotation:pdf', args=[plain.pk]) + query).content)).pages)
            self.assertNotIn('DISCOUNT', text)
            self.assertNotIn('AMOUNT IN WORDS', text)          # the words line has no label ...
            self.assertNotIn('DIRHAMS', text)                  # ... and no currency name
            order = ['TOTAL AMOUNT EXCL. VAT (AED)', '1,055.00', 'VAT 5% (AED)', '52.75', 'GRAND TOTAL (AED)', '1,107.75',
                     'AED : ONE THOUSAND ONE HUNDRED AND SEVEN AND SEVENTY FIVE FILS ONLY',
                     'TERMS AND CONDITIONS']
            positions = [text.index(x) for x in order]
            self.assertEqual(positions, sorted(positions), query)
        # with a 5% discount (52.75): total, discount, excl. VAT, VAT, grand total, words
        self.client.post(reverse('quotation:new'), self._payload(discount_type='percent', discount_value='5'))
        q = Quotation.objects.order_by('-id').first()
        self.assertEqual((q.discount_amount, q.net_total, q.vat_amount, q.grand_total),
                         (Decimal('52.75'), Decimal('1002.25'), Decimal('50.11'), Decimal('1052.36')))   # 50.1125 -> 50.11
        text = ' '.join(' '.join(p.extract_text() for p in PdfReader(BytesIO(
            self.client.get(reverse('quotation:pdf', args=[q.pk])).content)).pages).split())   # whitespace-normalised
        order = ['TOTAL (AED) 1,055.00', 'DISCOUNT 5% (AED) - 52.75', 'TOTAL AMOUNT EXCL. VAT (AED) 1,002.25',
                 'VAT 5% (AED) 50.11', 'GRAND TOTAL (AED) 1,052.36',
                 'AED : ONE THOUSAND AND FIFTY TWO AND THIRTY SIX FILS ONLY']
        positions = [text.index(x) for x in order]
        self.assertEqual(positions, sorted(positions))
        # an amount discount is labelled without a percentage
        self.client.post(reverse('quotation:new'), self._payload(discount_type='amount', discount_value='55'))
        amount_q = Quotation.objects.order_by('-id').first()
        text = ' '.join(' '.join(p.extract_text() for p in PdfReader(BytesIO(
            self.client.get(reverse('quotation:pdf', args=[amount_q.pk])).content)).pages).split())
        self.assertIn('DISCOUNT (AED) - 55.00', text)
        self.assertIn('AED : ONE THOUSAND AND FIFTY ONLY', text)

    def test_brand_and_color_are_always_capitals(self):
        # saved in capitals whatever was typed ...
        self.client.post(reverse('quotation:new'), self._payload(
            item_brand=['ariston Plus', 'pegler'], item_color=['brushed ss304', 'matt Black']))
        q = Quotation.objects.get()
        first, second = q.items.all()
        self.assertEqual((first.brand, first.color), ('ARISTON PLUS', 'BRUSHED SS304'))
        self.assertEqual((second.brand, second.color), ('PEGLER', 'MATT BLACK'))
        self.assertEqual(first.description, 'Heater 50L')                       # other columns are left as typed
        # ... printed in capitals even for a row that was saved in lowercase before this rule existed
        legacy = Quotation.objects.create(qtn_no='LEGACY-1', date='2026-10-20')
        QuotationItem.objects.create(quotation=legacy, position=0, brand='jomix', color='brushed ss304', size='820*500*145',
                                     description='mixed Case description', qty=Decimal('1'), unit_price=Decimal('10'))
        for query in ('', '?orientation=landscape'):
            text = ' '.join(PdfReader(BytesIO(self.client.get(reverse('quotation:pdf', args=[legacy.pk]) + query).content)).pages[0].extract_text().split())
            self.assertIn('JOMIX', text)
            self.assertIn('BRUSHED SS304', text)
            self.assertNotIn('jomix', text)
            self.assertNotIn('brushed ss304', text)
            self.assertIn('mixed Case description', text)                       # the description keeps its case
        # the edit form shows them in capitals and the inputs convert as you type
        item = legacy.items.get()
        page = self.client.get(reverse('quotation:edit', args=[legacy.pk]))
        self.assertContains(page, 'name="item_brand" class="input-x caps"')
        self.assertContains(page, 'name="item_color" class="input-x caps"')
        self.assertContains(page, "(d.brand || '').toUpperCase()")
        self.assertContains(page, 'toUpperCase()')

    def test_the_configured_stamp_users_in_settings(self):
        from importlib import import_module
        configured = import_module('salesorder.settings').QUOTATION_STAMP_USERS
        self.assertEqual(sorted(configured), ['manager', 'md', 'so'])
        for username, allowed in (('so', True), ('Manager', True), ('MD', True), ('soadmin', False), ('aman', False)):
            user = get_user_model().objects.create_user(username, password='x')
            with override_settings(QUOTATION_STAMP_USERS=configured):
                self.client.force_login(user)
                self.assertEqual(self.client.get(reverse('quotation:stamp_settings')).status_code,
                                 200 if allowed else 403, username)

    def test_other_apps_untouched(self):
        # The new app must not interfere with the Bangladesh PI routes.
        self.assertEqual(self.client.get(reverse('bangladeshpi:list')).status_code, 200)


def xlsx_file(rows, name='products.xlsx'):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type='application/octet-stream')


@override_settings(QUOTATION_PRIVATE_ROOT=TMP_ROOT, QUOTATION_PUBLIC_BASE_URL='https://example.test',
                   QUOTATION_STAMP_USERS=())
class ProductLibraryTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('libuser', password='x')
        self.client.force_login(self.user)

    def _excel(self, rows):
        return self.client.post(reverse('quotation:product_import_excel'), {'excel': xlsx_file(rows)})

    def _payload(self, **extra):
        return QuotationTests('_payload')._payload(**extra)

    def test_excel_import_creates_updates_and_never_wipes(self):
        head = ['Model No', 'Brand', 'Description', 'Photo file', 'Datasheet file']
        self._excel([head, ['H12', 'jomix', 'Basin mixer', '', ''], ['ab/12', 'x', 'Hose', 'IMG_1.jpg', ''],
                     [None, None, None, None, None], ['h12', 'a', 'dup', '', '']])
        self.assertEqual(Product.objects.count(), 2)
        h12 = Product.objects.get(model_key='H12')
        self.assertEqual((h12.brand, h12.description), ('JOMIX', 'Basin mixer'))
        self.assertEqual(Product.objects.get(model_key='AB-12').photo_name, 'IMG_1.jpg')
        self._excel([head, ['H12', '', 'New text', '', '']])   # blank brand keeps JOMIX; description updates
        h12.refresh_from_db()
        self.assertEqual((h12.brand, h12.description), ('JOMIX', 'New text'))
        self.assertEqual(Product.objects.count(), 2)

    def test_bad_excel_is_rejected(self):
        resp = self.client.post(reverse('quotation:product_import_excel'),
                                {'excel': SimpleUploadedFile('x.xlsx', b'junk')}, follow=True)
        self.assertContains(resp, 'could not be read')
        self.assertEqual(Product.objects.count(), 0)

    def test_files_attach_by_name_and_by_hint_and_report_unmatched(self):
        Product.objects.create(model_no='H12')
        Product.objects.create(model_no='AB/12')
        Product.objects.create(model_no='X1', photo_name='IMG_9.png')
        url = reverse('quotation:product_import_files')
        resp = self.client.post(url, {'kind': 'photo', 'files': [
            png_file('h12.PNG'), png_file('AB-12.png'), png_file('IMG_9.png'), png_file('ZZ.png')]})
        res = {r['name']: r['status'] for r in resp.json()['results']}
        self.assertEqual(res, {'h12.PNG': 'ok', 'AB-12.png': 'ok', 'IMG_9.png': 'ok', 'ZZ.png': 'unmatched'})
        self.assertTrue(all(Product.objects.get(model_key=k).photo for k in ('H12', 'AB-12', 'X1')))
        resp = self.client.post(url, {'kind': 'datasheet', 'files': [SimpleUploadedFile('H12.pdf', PDF_BYTES, 'application/pdf')]})
        self.assertEqual(resp.json()['results'][0]['status'], 'ok')
        self.assertEqual(jpeg_size(Product.objects.get(model_key='H12').datasheet.open('rb').read())[0], 1240)
        bad = self.client.post(url, {'kind': 'datasheet', 'files': [SimpleUploadedFile('H12.pdf', b'nope')]})
        self.assertEqual(bad.json()['results'][0]['status'], 'error')
        self.assertTrue(Product.objects.get(model_key='H12').datasheet)   # a bad file never removes the good one

    def test_model_no_is_unique_in_the_form(self):
        Product.objects.create(model_no='H12')
        resp = self.client.post(reverse('quotation:product_new'), {'model_no': ' h12 ', 'brand': 'a'})
        self.assertContains(resp, 'already exists')
        self.assertEqual(Product.objects.count(), 1)
        self.client.post(reverse('quotation:product_new'),
                         {'model_no': 'H13', 'brand': 'grohe', 'description': 'd', 'photo': png_file()})
        p = Product.objects.get(model_key='H13')
        self.assertEqual(p.brand, 'GROHE')
        self.assertTrue(p.photo)

    def test_search_prefers_models_starting_with_the_text(self):
        for m in ('XH1', 'H12', 'H13'):
            Product.objects.create(model_no=m, brand='B')
        names = [r['model_no'] for r in self.client.get(reverse('quotation:product_search') + '?q=h1').json()['results']]
        self.assertEqual(names, ['H12', 'H13', 'XH1'])
        self.assertEqual(self.client.get(reverse('quotation:product_search') + '?q=').json()['results'], [])

    def test_search_ignores_case_spaces_dashes_and_finds_descriptions(self):
        Product.objects.create(model_no='AB/12', brand='Grohe', description='Chrome basin MIXER')
        Product.objects.create(model_no='ZX 900', brand='Other', description='x')
        url = reverse('quotation:product_search')

        def find(q):
            return [r['model_no'] for r in self.client.get(url, {'q': q}).json()['results']]
        for q in ('ab-12', 'AB 12', 'ab/12', 'Ab12', 'grohe', 'basin mixer', 'ab'):
            self.assertEqual(find(q), ['AB/12'], q)
        self.assertEqual(find('zx900'), ['ZX 900'])
        self.assertEqual(find('zx 9'), ['ZX 900'])
        self.assertEqual(find('nothing-here'), [])

    def test_quotation_row_gets_its_own_copy_unaffected_by_library_changes(self):
        p = Product.objects.create(model_no='H12', brand='JOMIX', description='Mixer')
        url = reverse('quotation:product_import_files')
        self.client.post(url, {'kind': 'photo', 'files': [png_file('H12.png')]})
        self.client.post(url, {'kind': 'datasheet', 'files': [SimpleUploadedFile('H12.pdf', PDF_BYTES, 'application/pdf')]})
        p.refresh_from_db()
        self.client.post(reverse('quotation:new'), self._payload(lib_id=[str(p.pk), '']))
        item = QuotationItem.objects.get(brand='ARISTON')
        self.assertTrue(item.photo and item.datasheet and item.share_token)
        self.assertNotEqual(item.photo.name, p.photo.name)
        self.assertNotEqual(item.datasheet.name, p.datasheet.name)
        token, ds_bytes = item.share_token, item.datasheet.open('rb').read()
        # change, then delete the product: the quotation and the datasheet link already sent stay as they were
        self.client.post(url, {'kind': 'datasheet', 'files': [SimpleUploadedFile('H12.pdf', make_pdf('Other', 3), 'application/pdf')]})
        self.client.post(reverse('quotation:product_delete', args=[p.pk]))
        item.refresh_from_db()
        self.assertEqual(item.datasheet.open('rb').read(), ds_bytes)
        self.client.logout()
        self.assertEqual(self.client.get(reverse('quotation:datasheet', args=[token])).status_code, 200)

    def test_row_removal_beats_library(self):
        p = Product.objects.create(model_no='H12')
        self.client.post(reverse('quotation:product_import_files'), {'kind': 'photo', 'files': [png_file('H12.png')]})
        self.client.post(reverse('quotation:new'), self._payload(lib_id=[str(p.pk), ''], remove_photo=['1', '0']))
        self.assertFalse(QuotationItem.objects.get(brand='ARISTON').photo)

    def test_datasheet_can_be_viewed_by_logged_in_users_only(self):
        p = Product.objects.create(model_no='H12')
        url = reverse('quotation:product_datasheet', args=[p.pk])
        self.assertEqual(self.client.get(url).status_code, 404)   # no datasheet yet
        self.client.post(reverse('quotation:product_import_files'),
                         {'kind': 'datasheet', 'files': [SimpleUploadedFile('H12.pdf', PDF_BYTES, 'application/pdf')]})
        resp = self.client.get(url)
        self.assertEqual(resp['Content-Type'], 'image/jpeg')
        self.assertTrue(resp['Content-Disposition'].startswith('inline'))
        self.assertEqual(jpeg_size(b''.join(resp.streaming_content))[0], 1240)
        self.assertTrue(self.client.get(url + '?download=1')['Content-Disposition'].startswith('attachment'))
        self.assertContains(self.client.get(reverse('quotation:product_edit', args=[p.pk])), 'View datasheet')
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    def test_brand_filter_lists_every_brand_and_filters(self):
        Product.objects.create(model_no='A1', brand='GROHE')
        Product.objects.create(model_no='A2', brand='GROHE')
        Product.objects.create(model_no='B1', brand='JOMIX')
        Product.objects.create(model_no='C1')
        url = reverse('quotation:product_list')
        resp = self.client.get(url)
        self.assertEqual([b for b, _ in resp.context['brand_options']], ['GROHE', 'JOMIX', '__none__'])   # each brand once
        self.assertEqual(len(resp.context['page']), 4)
        self.assertContains(resp, 'Select all')
        self.assertContains(resp, 'Clear all')

        def shown(**kw):
            return sorted(p.model_no for p in self.client.get(url, kw).context['page'])
        self.assertEqual(shown(brand=['GROHE']), ['A1', 'A2'])
        self.assertEqual(shown(brand=['grohe']), ['A1', 'A2'])
        self.assertEqual(shown(brand=['GROHE', 'JOMIX']), ['A1', 'A2', 'B1'])      # several brands at once
        self.assertEqual(shown(brand=['JOMIX'], q='b'), ['B1'])
        self.assertEqual(shown(brand=['__none__']), ['C1'])
        self.assertEqual(shown(brand=['GROHE', 'JOMIX', '__none__']), ['A1', 'A2', 'B1', 'C1'])   # Select all = everything
        self.assertEqual(shown(), ['A1', 'A2', 'B1', 'C1'])                                        # Clear all = everything
        Product.objects.create(model_no='D1', brand='HANSGROHE')        # a new brand appears by itself
        self.assertIn('HANSGROHE', [b for b, _ in self.client.get(url).context['brand_options']])

    def test_customer_suggestions_from_earlier_quotations(self):
        def make(date, name, address, trn='', contact=''):
            self.client.post(reverse('quotation:new'), self._payload(
                date=date, customer_name=name, customer_address=address, customer_trn=trn, customer_contact=contact))
        make('2026-03-01', 'Al Noor Trading LLC', 'Old address', '111', '0501')
        make('2026-09-01', 'AL NOOR  trading llc', 'New address\nDubai', '222', '0502')   # same customer, newer
        make('2026-05-01', 'Beta Contracting', 'Sharjah')
        url = reverse('quotation:customer_search')

        def find(q):
            return self.client.get(url, {'q': q}).json()['results']
        res = find('al noor')
        self.assertEqual(len(res), 1)   # one customer, not two
        self.assertEqual((res[0]['address'], res[0]['trn'], res[0]['contact']), ('New address\nDubai', '222', '0502'))   # newest quotation
        self.assertEqual(find('NOOR trading')[0]['name'], 'AL NOOR trading llc')
        self.assertEqual(find('zzz'), [])
        self.assertEqual(find(''), [])
        self.client.logout()
        self.assertEqual(self.client.get(url, {'q': 'al'}).status_code, 302)

    def test_library_pages_need_login(self):
        self.client.logout()
        for name in ('product_list', 'product_new', 'product_search', 'product_template'):
            self.assertEqual(self.client.get(reverse(f'quotation:{name}')).status_code, 302, name)
        self.assertEqual(self.client.post(reverse('quotation:product_import_files'), {'kind': 'photo'}).status_code, 302)

    def test_pages_render_and_template_downloads(self):
        Product.objects.create(model_no='H12', brand='JOMIX')
        self.assertContains(self.client.get(reverse('quotation:product_list')), 'H12')
        self.assertEqual(self.client.get(reverse('quotation:product_edit', args=[Product.objects.get().pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('quotation:product_template')).status_code, 200)
        self.assertEqual(self.client.get(reverse('quotation:new')).status_code, 200)


@override_settings(QUOTATION_PRIVATE_ROOT=TMP_ROOT, QUOTATION_PUBLIC_BASE_URL='https://example.test', QUOTATION_STAMP_USERS=(),
                   QUOTATION_EMAIL_HOST='smtp.test', QUOTATION_EMAIL_USER='sanitarysales@junaid.ae', QUOTATION_EMAIL_PASSWORD='secret')
class QuotationEmailTests(TestCase):
    def setUp(self):
        from unittest import mock
        self.user = get_user_model().objects.create_user('mailer', password='x')
        self.client.force_login(self.user)
        self.client.post(reverse('quotation:new'), QuotationTests('_payload')._payload(customer_name='Acme'))
        self.q = Quotation.objects.get()
        self.url = reverse('quotation:email', args=[self.q.pk])
        patcher = mock.patch('quotation.mailer._save_copy_to_sent', return_value=True)   # never touch the real IMAP server
        self.copy = patcher.start()
        self.addCleanup(patcher.stop)

    def _post(self, **kw):
        data = {'to': 'buyer@acme.com', 'cc': 'boss@acme.com', 'subject': 'Quotation', 'body': 'Hello', 'orientation': 'portrait'}
        data.update(kw)
        return self.client.post(self.url, data)

    def test_sends_pdf_from_the_quotation_mailbox_and_logs_it(self):
        from django.core import mail
        resp = self._post()
        self.assertRedirects(resp, reverse('quotation:list'))
        self.assertEqual(len(mail.outbox), 1)
        m = mail.outbox[0]
        self.assertIn('sanitarysales@junaid.ae', m.from_email)
        self.assertEqual((m.to, m.cc, m.reply_to), (['buyer@acme.com'], ['boss@acme.com'], ['sanitarysales@junaid.ae']))
        name, content, mimetype = m.attachments[0]
        self.assertEqual((name, mimetype), (f'Quotation_{self.q.qtn_no}.pdf', 'application/pdf'))
        self.assertEqual(content[:5], b'%PDF-')
        log = QuotationEmail.objects.get()
        self.assertTrue(log.success and log.saved_to_sent)
        self.assertEqual((log.to_addrs, log.sent_by), ('buyer@acme.com', self.user))
        self.q.refresh_from_db()
        self.assertEqual(self.q.customer_email, 'buyer@acme.com')   # remembered for next time
        self.assertContains(self.client.get(reverse('quotation:list')), 'Mail sent')
        self.assertContains(self.client.get(self.url), 'buyer@acme.com')   # the sent-mails section

    def test_email_always_carries_the_stamp_without_changing_the_quotation(self):
        from unittest import mock
        with mock.patch('quotation.email_views.build_quotation_pdf', return_value=b'%PDF-1.4 x') as build:
            self._post()
        quotation, _, _ = build.call_args.args[:3]
        self.assertTrue(quotation.add_stamp)
        self.assertTrue(build.call_args.kwargs['with_stamp'])
        self.q.refresh_from_db()
        self.assertFalse(self.q.add_stamp)   # the saved "Add stamp" tick is untouched

    def test_failure_is_logged_and_shown(self):
        from unittest import mock
        with mock.patch('django.core.mail.EmailMessage.send', side_effect=OSError('connection refused')):
            resp = self._post()
        self.assertContains(resp, 'connection refused')
        log = QuotationEmail.objects.get()
        self.assertFalse(log.success)
        self.assertIn('connection refused', log.error)
        self.q.refresh_from_db()
        self.assertEqual(self.q.customer_email, '')   # nothing remembered after a failure
        self.assertContains(self.client.get(reverse('quotation:list')), 'Mail failed')

    def test_validation_and_not_configured(self):
        from django.core import mail
        self.assertContains(self._post(to='not-an-email'), 'valid email')
        self.assertContains(self._post(to=''), 'customer email')
        self.assertEqual(len(mail.outbox), 0)
        with override_settings(QUOTATION_EMAIL_PASSWORD=''):
            self.assertContains(self._post(), 'not set up')
        self.assertEqual((len(mail.outbox), QuotationEmail.objects.count()), (0, 0))

    def test_several_recipients_and_login_required(self):
        from django.core import mail
        self._post(to='a@x.com; b@y.com, a@x.com')
        self.assertEqual(mail.outbox[0].to, ['a@x.com', 'b@y.com'])
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.assertEqual(self.client.post(self.url, {'to': 'a@x.com'}).status_code, 302)

    def test_customer_email_saved_with_quotation_and_suggested(self):
        self.client.post(reverse('quotation:edit', args=[self.q.pk]), QuotationTests('_payload')._payload(
            customer_name='Acme', customer_email='buyer@acme.com', item_id=[str(i.pk) for i in self.q.items.all()]))
        self.q.refresh_from_db()
        self.assertEqual(self.q.customer_email, 'buyer@acme.com')
        self.assertEqual(self.client.get(reverse('quotation:customer_search'), {'q': 'acme'}).json()['results'][0]['email'], 'buyer@acme.com')


@override_settings(QUOTATION_PRIVATE_ROOT=TMP_ROOT, QUOTATION_PUBLIC_BASE_URL='https://example.test')
class CustomerIdEmailTests(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_user('cid', password='x'))

    def test_saved_suggested_and_printed(self):
        from unittest import mock
        self.client.post(reverse('quotation:new'), QuotationTests('_payload')._payload(
            customer_name='Acme', customer_code='C00123', customer_email='buyer@acme.com', customer_person='Mr Zed'))
        q = Quotation.objects.get()
        self.assertEqual((q.customer_code, q.customer_email), ('C00123', 'buyer@acme.com'))
        res = self.client.get(reverse('quotation:customer_search'), {'q': 'acme'}).json()['results'][0]
        self.assertEqual((res['code'], res['email'], res['person']), ('C00123', 'buyer@acme.com', 'Mr Zed'))
        self.assertContains(self.client.get(reverse('quotation:edit', args=[q.pk])), 'C00123')
        from quotation.pdf_builder import _header_block
        with mock.patch('quotation.pdf_builder.Paragraph', wraps=__import__('reportlab.platypus', fromlist=['Paragraph']).Paragraph) as p:
            resp = self.client.get(reverse('quotation:pdf', args=[q.pk]))
        self.assertEqual(resp.status_code, 200)
        texts = ' '.join(str(c.args[0]) for c in p.call_args_list if c.args)
        self.assertIn('C00123', texts)
        self.assertIn('Mr Zed', texts)
        self.assertIn('buyer@acme.com', texts)


@override_settings(QUOTATION_PRIVATE_ROOT=TMP_ROOT)
class CustomerListTests(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_user('cl', password='x'))

    def _xlsx(self, rows, header=('Customer ID', 'Customer Name', 'TRN Number', 'Contact Number', 'Contact Person', 'Address', 'Customer Email')):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(list(header))
        for r in rows:
            ws.append(list(r))
        buf = BytesIO()
        wb.save(buf)
        return SimpleUploadedFile('c.xlsx', buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def _import(self, rows, **kw):
        return self.client.post(reverse('quotation:customer_import_excel'), {'excel': self._xlsx(rows, **kw)}, follow=True)

    def test_import_adds_updates_and_never_wipes(self):
        self._import([('V00001', 'SUNJET TRDG. CO. LLC.', None, None, None, None, None),
                      ('V00002', 'ACME LLC', 100123456700003, '+971501234567', 'Mr Ahmed', 'PO BOX 1\nDubai', 'a@acme.com')])
        self.assertEqual(Customer.objects.count(), 2)
        c = Customer.objects.get(code='V00002')
        self.assertEqual((c.trn, c.contact_number, c.contact_person, c.email), ('100123456700003', '+971501234567', 'Mr Ahmed', 'a@acme.com'))
        resp = self._import([('v00002', 'ACME LLC NEW', None, None, None, None, None), ('V00001', 'x', '', '', '', '', ''),
                             ('V00001', 'dup', '', '', '', '', ''), ('', 'No id', '', '', '', '', ''), ('V9', None, '', '', '', '', '')])
        self.assertEqual(Customer.objects.count(), 2)
        c.refresh_from_db()
        self.assertEqual((c.name, c.trn, c.email), ('ACME LLC NEW', '100123456700003', 'a@acme.com'))   # blank cells kept the old data
        self.assertContains(resp, 'repeated in the sheet')
        self.assertContains(resp, 'Customer ID is empty')
        self.assertContains(resp, 'no Customer Name')

    def test_bad_files_are_refused(self):
        resp = self.client.post(reverse('quotation:customer_import_excel'), {'excel': SimpleUploadedFile('c.txt', b'x')}, follow=True)
        self.assertContains(resp, 'Please upload an .xlsx')
        resp = self._import([('V1', 'A')], header=('Foo', 'Bar'))
        self.assertContains(resp, 'must have')
        self.assertEqual(Customer.objects.count(), 0)

    def test_add_edit_delete_and_unique_id(self):
        self.client.post(reverse('quotation:customer_new'), {'code': 'C1', 'name': 'Acme', 'email': 'a@x.com'})
        c = Customer.objects.get()
        resp = self.client.post(reverse('quotation:customer_new'), {'code': 'c1', 'name': 'Other'})
        self.assertContains(resp, 'already exists')
        self.client.post(reverse('quotation:customer_edit', args=[c.pk]), {'code': 'C1', 'name': 'Acme 2'})
        c.refresh_from_db()
        self.assertEqual(c.name, 'Acme 2')
        self.assertContains(self.client.get(reverse('quotation:customer_list'), {'q': 'acme'}), 'Acme 2')
        self.client.post(reverse('quotation:customer_delete', args=[c.pk]))
        self.assertEqual(Customer.objects.count(), 0)

    def test_login_required_and_template(self):
        self.assertEqual(self.client.get(reverse('quotation:customer_template')).status_code, 200)
        self.client.logout()
        for name in ('customer_list', 'customer_new', 'customer_template'):
            self.assertEqual(self.client.get(reverse(f'quotation:{name}')).status_code, 302)
        self.assertEqual(self.client.post(reverse('quotation:customer_import_excel')).status_code, 302)

    def test_listed_customer_fills_the_form_and_wins_over_old_quotations(self):
        self.client.post(reverse('quotation:new'), QuotationTests('_payload')._payload(
            customer_name='Acme Trading', customer_address='Old address', customer_trn='111'))
        Customer.objects.create(code='V5', name='ACME TRADING', trn='222', contact_number='050', contact_person='Mr Z',
                                address='New address', email='n@acme.com')
        Customer.objects.create(code='V6', name='Zeta Only', trn='333')
        res = self.client.get(reverse('quotation:customer_search'), {'q': 'acme'}).json()['results']
        self.assertEqual(len(res), 1)
        self.assertEqual((res[0]['code'], res[0]['trn'], res[0]['address'], res[0]['person'], res[0]['contact'], res[0]['email']),
                         ('V5', '222', 'New address', 'Mr Z', '050', 'n@acme.com'))
        self.assertEqual(self.client.get(reverse('quotation:customer_search'), {'q': 'zeta'}).json()['results'][0]['code'], 'V6')

    def test_editing_a_customer_does_not_change_a_made_quotation(self):
        self.client.post(reverse('quotation:new'), QuotationTests('_payload')._payload(customer_name='Acme', customer_code='V1', customer_trn='111'))
        Customer.objects.create(code='V1', name='Acme', trn='999')
        Customer.objects.all().delete()
        q = Quotation.objects.get()
        self.assertEqual((q.customer_code, q.customer_trn), ('V1', '111'))


@override_settings(QUOTATION_PRIVATE_ROOT=TMP_ROOT)
class ProductPricingTests(TestCase):
    def setUp(self):
        from so.models import Items, Role
        self.Items = Items
        self.admin = get_user_model().objects.create_user('boss', password='x')
        Role.objects.create(user=self.admin, role='Admin')
        self.staff = get_user_model().objects.create_user('sales', password='x')
        Role.objects.create(user=self.staff, role='Salesman')
        Items.objects.create(item_code='100001', item_description='Heater', item_upvc='ar-50', item_cost=100.0, item_firm='ARISTON')
        Items.objects.create(item_code='100002', item_description='No cost', item_upvc='NC-1', item_cost=0.0, item_firm='X')
        self.client.force_login(self.admin)

    def _new(self, model_no='AR 50', **extra):
        data = {'model_no': model_no, 'brand': 'ariston', 'description': 'Heater'}
        data.update(extra)
        self.client.post(reverse('quotation:product_new'), data)
        return Product.objects.get(model_key=product_key(model_no))

    def _edit(self, p, **extra):
        data = {'model_no': p.model_no, 'brand': p.brand, 'description': p.description}
        data.update(extra)
        return self.client.post(reverse('quotation:product_edit', args=[p.pk]), data)

    def test_new_product_gets_stock_cost_by_upc_and_double_price(self):
        from decimal import Decimal as D
        p = self._new()                       # 'AR 50' finds the stock UPC code 'ar-50' (case, space, dash ignored)
        self.assertEqual((p.cost, p.selling_price, p.cost_manual, p.price_manual), (D('100.00'), D('200.00'), False, False))
        self.assertIsNone(self._new('NC-1').cost)       # a 0 cost in the stock data means unknown
        self.assertIsNone(self._new('UNKNOWN-9').cost)  # no stock item with that UPC code

    def test_admin_edits_are_kept_by_refresh_and_reset_restores(self):
        from decimal import Decimal as D
        p = self._new()
        self._edit(p, cost='120', selling_price='240.00')   # a cost typed by hand; the price stays automatic (2 x 120)
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price, p.cost_manual, p.price_manual), (D('120.00'), D('240.00'), True, False))
        self.assertRedirects(self.client.post(reverse('quotation:product_refresh_costs')), reverse('quotation:product_list'))
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price), (D('120.00'), D('240.00')))   # the refresh did not touch the hand-set cost
        self._edit(p, cost='120', selling_price='999')       # a price typed by hand
        p.refresh_from_db()
        self.assertEqual((p.selling_price, p.price_manual), (D('999.00'), True))
        self._edit(p, cost='130', selling_price='999')       # the cost moves, the hand-set price stays
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price), (D('130.00'), D('999.00')))
        self._edit(p, cost='130', selling_price='999', reset_cost='1', reset_price='1')
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price, p.cost_manual, p.price_manual), (D('100.00'), D('200.00'), False, False))

    def test_refresh_updates_automatic_costs_only(self):
        from decimal import Decimal as D
        p = self._new()
        self.Items.objects.filter(item_code='100001').update(item_cost=150.0)
        self.client.post(reverse('quotation:product_refresh_costs'))
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price), (D('150.00'), D('300.00')))

    def test_retyping_the_automatic_value_is_not_a_manual_override(self):
        from decimal import Decimal as D
        p = self._new()
        self._edit(p, cost='100', selling_price='200')
        p.refresh_from_db()
        self.assertEqual((p.cost_manual, p.price_manual), (False, False))
        self.Items.objects.filter(item_code='100001').update(item_cost=110.0)
        self.client.post(reverse('quotation:product_refresh_costs'))
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price), (D('110.00'), D('220.00')))

    def test_invalid_amounts_are_refused(self):
        p = self._new()
        resp = self._edit(p, cost='abc')
        self.assertContains(resp, 'Cost must be a valid amount')
        resp = self._edit(p, cost='100', selling_price='-5')
        self.assertContains(resp, 'Selling price must be a valid amount')

    def test_cost_is_hidden_from_everyone_but_admin(self):
        p = self._new()
        self.assertContains(self.client.get(reverse('quotation:product_list')), '<th>Cost</th>')
        self.client.force_login(self.staff)
        html = self.client.get(reverse('quotation:product_list')).content.decode()
        self.assertNotIn('<th>Cost</th>', html)
        self.assertNotIn('100.00', html)
        self.assertIn('200.00', html)                               # the selling price is visible
        form = self.client.get(reverse('quotation:product_edit', args=[p.pk])).content.decode()
        self.assertNotIn('name="cost"', form)
        self.assertNotIn('100.00', form)
        res = self.client.get(reverse('quotation:product_search'), {'q': 'ar'}).json()['results'][0]
        self.assertEqual(res['selling_price'], '200.00')
        self.assertNotIn('cost', res)
        self.assertEqual(self.client.post(reverse('quotation:product_refresh_costs')).status_code, 403)
        self._edit(p, cost='1', selling_price='2')                    # posted prices from a non-admin are ignored
        p.refresh_from_db()
        self.assertEqual((p.cost, p.selling_price), (__import__('decimal').Decimal('100.00'), __import__('decimal').Decimal('200.00')))

    def test_excel_import_fills_cost_for_new_products(self):
        from decimal import Decimal as D
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(['Model No', 'Brand'])
        wb.active.append(['AR-50', 'Ariston'])
        buf = BytesIO()
        wb.save(buf)
        self.client.post(reverse('quotation:product_import_excel'), {'excel': SimpleUploadedFile('p.xlsx', buf.getvalue())})
        p = Product.objects.get()
        self.assertEqual((p.cost, p.selling_price), (D('100.00'), D('200.00')))
