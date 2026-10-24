"""Product library: add products once (model no, brand, description, photo, datasheet), then pick them in a quotation."""
import logging
import os
import re
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Q, Value
from django.db.models.functions import Replace
from django.http import FileResponse, Http404, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_GET, require_POST

from .datasheet import DatasheetError, to_jpeg
from . import pricing
from .models import Product, product_key
from .views import MAX_DATASHEET_BYTES, _check_photo, _display_name

log = logging.getLogger('quotation')

NO_BRAND = '__none__'   # brand-filter value for products without a brand
MAX_EXCEL_BYTES =5 * 1024 * 1024
MAX_EXCEL_ROWS = 5000
TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), 'sample_files', 'product_library_template.xlsx')
HEADERS = {
    'model no': 'model_no', 'model': 'model_no', 'model number': 'model_no',
    'brand': 'brand', 'description': 'description',
    'photo file': 'photo_name', 'datasheet file': 'datasheet_name',
}


def _back_url(request):
    """Where to return to after saving / deleting: the Products list as the user left it (brand filter, search, page).
    Only an address on this site that points into the Products list is accepted; anything else gives the plain list."""
    plain = reverse('quotation:product_list')
    nxt = (request.POST.get('next') or request.GET.get('next') or '').strip()
    if nxt.startswith(plain) and url_has_allowed_host_and_scheme(nxt, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        return nxt
    return plain


def _delete_files(product):
    if product.photo:
        product.photo.delete(save=False)
    if product.datasheet:
        product.datasheet.delete(save=False)


def _set_photo(product, f):
    if product.photo:
        product.photo.delete(save=False)
    product.photo.save(f.name, f, save=False)


def _set_datasheet(product, name, data):
    from django.core.files.base import ContentFile
    if product.datasheet:
        product.datasheet.delete(save=False)
    product.datasheet.save(name, ContentFile(data), save=False)


def _filtered(request):
    """The products matching the search box and the brand filter (shared by the list and its exports)."""
    q = (request.GET.get('q') or '').strip()
    products = Product.objects.all()
    if q:
        products = products.filter(Q(model_key__icontains=product_key(q)) | Q(brand__icontains=q) | Q(description__icontains=q))
    # Brand filter: the choices are read from the products themselves, so a new brand shows up as soon as a product has it.
    brands = list(Product.objects.exclude(brand='').order_by('brand').values_list('brand', flat=True).distinct())
    chosen = [b.strip() for b in request.GET.getlist('brand') if b.strip()]
    chosen_low = {b.lower() for b in chosen}
    # every option ticked (Select all) or none ticked (Clear all) means no filtering
    if chosen_low and chosen_low >= {b.lower() for b in brands} | {NO_BRAND}:
        chosen, chosen_low = [], set()
    if chosen_low:
        cond = Q()
        for b in chosen:
            cond |= Q(brand='') if b == NO_BRAND else Q(brand__iexact=b)
        products = products.filter(cond)
    return products, q, brands, chosen, chosen_low


@login_required
def product_list(request):
    products, q, brands, chosen, chosen_low = _filtered(request)
    page = Paginator(products, 50).get_page(request.GET.get('page'))
    return render(request, 'quotation/product_list.html', {
        'page': page, 'q': q, 'total': Product.objects.count(),
        'brand_options': [(b, b.lower() in chosen_low) for b in brands] + [(NO_BRAND, NO_BRAND in chosen_low)],
        'brand_count': len(chosen), 'brand_label': (chosen[0] if len(chosen) == 1 else f'{len(chosen)} brands') if chosen else 'All brands',
        'no_brand': NO_BRAND,
        'filter_qs': urlencode([('q', q)] + [('brand', b) for b in chosen]),
        'shown': page.paginator.count, 'is_admin': pricing.is_price_admin(request.user), 'here': request.get_full_path(),
        # lets the page check file names against the library before anything is uploaded
        'keys': {p.model_key: p.model_no for p in Product.objects.only('model_key', 'model_no')},
        'hints': {'photo': {product_key(p.photo_name): p.model_no for p in Product.objects.exclude(photo_name='')},
                  'datasheet': {product_key(p.datasheet_name): p.model_no for p in Product.objects.exclude(datasheet_name='')}},
    })


@login_required
def product_form(request, pk=None):
    product = get_object_or_404(Product, pk=pk) if pk else None
    errors = []
    is_admin = pricing.is_price_admin(request.user)
    values = {'model_no': product.model_no if product else '', 'brand': product.brand if product else '',
              'description': product.description if product else '',
              'cost': product.cost if product and product.cost is not None else '',
              'selling_price': product.selling_price if product and product.selling_price is not None else ''}
    if request.method == 'POST':
        values = {k: (request.POST.get(k) or '').strip() for k in ('model_no', 'brand', 'description')}
        values['brand'] = values['brand'].upper()
        new_cost = new_price = None
        if is_admin:   # only admin may touch cost / selling price; for anyone else these posted values are ignored
            values['cost'], values['selling_price'] = (request.POST.get('cost') or '').strip(), (request.POST.get('selling_price') or '').strip()
            try:
                new_cost = pricing.parse_money(values['cost'])
            except ValueError:
                errors.append('Cost must be a valid amount (0 or more).')
            try:
                new_price = pricing.parse_money(values['selling_price'])
            except ValueError:
                errors.append('Selling price must be a valid amount (0 or more).')
        photo, ds = request.FILES.get('photo'), request.FILES.get('datasheet')
        if not values['model_no']:
            errors.append('Model No is required.')
        elif Product.objects.filter(model_key=product_key(values['model_no'])).exclude(pk=product.pk if product else None).exists():
            errors.append(f'Model No "{values["model_no"]}" already exists in the library.')
        if photo:
            err = _check_photo(photo, 'Photo')
            if err:
                errors.append(err)
        ds_jpeg = None
        if ds:
            if ds.size > MAX_DATASHEET_BYTES:
                errors.append('Datasheet is too large (max 10 MB).')
            else:
                try:
                    ds_jpeg = to_jpeg(ds)
                except DatasheetError as e:
                    errors.append(f'Datasheet: {e}')
        if not errors:
            product = product or Product()
            model_changed = product.pk is None or product.model_no != values['model_no']
            product.model_no, product.brand, product.description = values['model_no'], values['brand'], values['description']
            _apply_pricing(request, product, is_admin, new_cost, new_price, model_changed)
            if request.POST.get('remove_photo') == '1' and product.photo:
                product.photo.delete(save=False)
            if photo:
                _set_photo(product, photo)
            if request.POST.get('remove_ds') == '1' and product.datasheet and not ds_jpeg:
                product.datasheet.delete(save=False)
            if ds_jpeg:
                _set_datasheet(product, *ds_jpeg)
            product.save()
            messages.success(request, f'Product {product.model_no} saved.')
            return redirect(_back_url(request))
    return render(request, 'quotation/product_form.html', {'product': product, 'values': values, 'errors': errors, 'is_admin': is_admin, 'back': _back_url(request)})


def _apply_pricing(request, product, is_admin, new_cost, new_price, model_changed):
    """Cost / selling price rules (see pricing.py). A typed value that differs from the automatic one is kept as manual."""
    if not is_admin:
        if model_changed and not product.cost_manual:   # a new / renamed product still gets its stock cost
            pricing.refresh_cost(product)
        return
    prev_cost, prev_price = product.cost, product.selling_price
    stock = pricing.stock_cost(product.model_no)
    if request.POST.get('reset_cost') == '1' or new_cost is None:
        product.cost, product.cost_manual = stock, False         # back to automatic (blank if the stock data has none)
    elif new_cost == prev_cost and product.cost_manual:
        pass                                                      # a hand-set cost, left as it is
    elif new_cost == stock:
        product.cost, product.cost_manual = new_cost, False
    elif new_cost == prev_cost:
        if model_changed and stock is not None:                   # automatic cost of a renamed product
            product.cost = stock
    else:
        product.cost, product.cost_manual = new_cost, True
    auto = pricing.auto_price(product.cost)
    if request.POST.get('reset_price') == '1' or new_price is None or new_price == auto:
        product.price_manual, product.selling_price = False, auto
    elif new_price == prev_price:
        product.selling_price = prev_price if product.price_manual else auto   # untouched
    else:
        product.price_manual, product.selling_price = True, new_price


@login_required
@require_POST
def product_refresh_costs(request):
    """Admin: refill every product's cost from the stock data. Costs set by hand are skipped."""
    if not pricing.is_price_admin(request.user):
        raise PermissionDenied
    cost_map = pricing.stock_cost_map()
    updated = unchanged = missing = manual = 0
    for product in Product.objects.all():
        if product.cost_manual:
            manual += 1
            continue
        old = product.cost
        if not pricing.refresh_cost(product, cost_map):
            missing += 1
            continue
        if product.cost == old:
            unchanged += 1
            continue
        product.save()
        updated += 1
    messages.success(request, f'Costs refreshed from stock: {updated} updated, {unchanged} unchanged, {missing} with no stock match, '
                              f'{manual} kept (set by hand).')
    return redirect(_back_url(request))


@login_required
@require_POST
def product_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)
    name = product.model_no
    _delete_files(product)
    product.delete()
    messages.success(request, f'Product {name} deleted. Quotations that already used it are not changed.')
    return redirect(_back_url(request))


@login_required
@require_GET
def product_photo(request, pk):
    """Login-only preview of a library photo (the files are private)."""
    product = get_object_or_404(Product, pk=pk)
    if not product.photo:
        raise Http404
    try:
        resp = FileResponse(product.photo.open('rb'))
    except FileNotFoundError:
        raise Http404
    resp['Cache-Control'] = 'private, max-age=300'
    return resp


@login_required
@require_GET
def product_datasheet(request, pk):
    """Login-only view of a library datasheet (a JPEG). Library datasheets have no public link; quotations get their own."""
    product = get_object_or_404(Product, pk=pk)
    if not product.datasheet:
        raise Http404
    try:
        resp = FileResponse(product.datasheet.open('rb'), content_type='image/jpeg')
    except FileNotFoundError:
        raise Http404
    resp['Content-Disposition'] = f'{"attachment" if request.GET.get("download") == "1" else "inline"}; filename="{_display_name(product.datasheet)}"'
    resp['Cache-Control'] = 'private, no-store'
    resp['X-Robots-Tag'] = 'noindex, nofollow'
    return resp


@login_required
@require_GET
def product_search(request):
    """Model No suggestions for the quotation form."""
    q = (request.GET.get('q') or '').strip()
    if not q:
        return JsonResponse({'results': []})
    # Not case sensitive, and spaces / dashes / slashes are ignored (ab 12, AB-12 and ab/12 find the same model).
    loose = re.sub(r'[\s\-/\\]', '', q).upper()
    products = list(Product.objects
                    .annotate(loose_key=Replace(Replace(Replace('model_key', Value('-'), Value('')), Value(' '), Value('')), Value('\t'), Value('')))
                    .filter(Q(loose_key__icontains=loose) | Q(brand__icontains=q) | Q(description__icontains=q))
                    .order_by('model_key')[:80])
    products.sort(key=lambda p: (not re.sub(r'[\s\-]', '', p.model_key).startswith(loose), p.model_key))   # models starting with the typed text first
    products = products[:30]
    return JsonResponse({'results': [{
        'id': p.pk, 'model_no': p.model_no, 'brand': p.brand, 'description': p.description,
        'photo_url': reverse('quotation:product_photo', args=[p.pk]) + f'?v={int(p.updated_at.timestamp())}' if p.photo else '',
        'datasheet_name': _display_name(p.datasheet) if p.datasheet else '',
        'selling_price': str(p.selling_price) if p.selling_price is not None else '',   # never the cost
    } for p in products]})


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


@login_required
@require_POST
def product_import_excel(request):
    """Step 1: read the Excel template and create / update products (text only; files come in step 2)."""
    f = request.FILES.get('excel')
    if not f:
        messages.error(request, 'Please choose the Excel file.')
        return redirect(_back_url(request))
    if f.size > MAX_EXCEL_BYTES or not f.name.lower().endswith('.xlsx'):
        messages.error(request, 'Please upload an .xlsx Excel file (max 5 MB).')
        return redirect(_back_url(request))
    from openpyxl import load_workbook
    try:
        sheet = load_workbook(f, read_only=True, data_only=True).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True, max_row=MAX_EXCEL_ROWS + 1))
    except Exception:
        messages.error(request, 'This file could not be read. Please use the downloaded template (.xlsx).')
        return redirect(_back_url(request))
    cols = {}
    for i, h in enumerate(rows[0] if rows else ()):
        field = HEADERS.get(_cell(h).lower())
        if field and field not in cols:
            cols[field] = i
    if 'model_no' not in cols:
        messages.error(request, 'The first row must have a "Model No" column. Please use the downloaded template.')
        return redirect(_back_url(request))

    created = updated = 0
    problems = []
    seen = set()
    cost_map = None   # read once, only if a new product turns up
    for n, row in enumerate(rows[1:], start=2):
        data = {k: _cell(row[i]) if i < len(row) else '' for k, i in cols.items()}
        model_no = data.get('model_no', '')
        if not model_no:
            if any(data.values()):
                problems.append(f'Row {n}: Model No is empty, skipped.')
            continue
        if len(model_no) > 100:
            problems.append(f'Row {n}: Model No is too long, skipped.')
            continue
        key = product_key(model_no)
        if key in seen:
            problems.append(f'Row {n}: {model_no} is repeated in the sheet, skipped.')
            continue
        seen.add(key)
        product = Product.objects.filter(model_key=key).first()
        is_new = product is None
        product = product or Product(model_no=model_no)
        # a blank cell never wipes what is already in the library
        if data.get('brand'):
            product.brand = data['brand'].upper()[:150]
        if data.get('description'):
            product.description = data['description']
        if data.get('photo_name'):
            product.photo_name = data['photo_name'][:200]
        if data.get('datasheet_name'):
            product.datasheet_name = data['datasheet_name'][:200]
        if is_new:   # a new product gets its cost from the stock data (and the selling price follows)
            cost_map = pricing.stock_cost_map() if cost_map is None else cost_map
            pricing.refresh_cost(product, cost_map)
        product.save()
        created += is_new
        updated += not is_new
    messages.success(request, f'Excel imported: {created} new, {updated} updated. Now upload the photos and datasheets below.')
    for p in problems[:10]:
        messages.warning(request, p)
    if len(problems) > 10:
        messages.warning(request, f'…and {len(problems) - 10} more rows skipped.')
    return redirect(_back_url(request))


@login_required
@require_POST
def product_import_files(request):
    """Step 2: photos or datasheets, matched to products by file name. Called in small batches by the page; returns JSON."""
    kind = request.POST.get('kind')
    if kind not in ('photo', 'datasheet'):
        return JsonResponse({'error': 'bad request'}, status=400)
    hint_field = 'photo_name' if kind == 'photo' else 'datasheet_name'
    hints = {product_key(getattr(p, hint_field)): p for p in Product.objects.exclude(**{hint_field: ''})}
    results = []
    for f in request.FILES.getlist('files'):
        stem = re.sub(r'\.[^.]*$', '', os.path.basename(f.name))
        product = hints.get(product_key(os.path.basename(f.name))) or Product.objects.filter(model_key=product_key(stem)).first()
        res = {'name': f.name, 'status': 'unmatched', 'model_no': '', 'message': 'No product with this Model No.'}
        if product:
            res['model_no'] = product.model_no
            try:
                if kind == 'photo':
                    err = _check_photo(f, 'Photo')
                    if err:
                        raise ValueError(err)
                    _set_photo(product, f)
                else:
                    if f.size > MAX_DATASHEET_BYTES:
                        raise ValueError('Datasheet is too large (max 10 MB).')
                    _set_datasheet(product, *to_jpeg(f))
                product.save()
                res.update(status='ok', message='Saved.')
            except (ValueError, DatasheetError) as e:
                res.update(status='error', message=str(e))
            except Exception:
                log.exception('Product library file import failed: %s', f.name)
                res.update(status='error', message='Could not save this file.')
        results.append(res)
    return JsonResponse({'results': results})


@login_required
@require_GET
def product_template(request):
    if not os.path.exists(TEMPLATE_PATH):
        raise Http404
    return FileResponse(open(TEMPLATE_PATH, 'rb'), as_attachment=True, filename='product_library_template.xlsx')
