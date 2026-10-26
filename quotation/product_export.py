"""Export the Products list (with the search / brand filter that is on screen) to Excel or PDF.

Cost is included only for admin users, exactly as on the Products page.
"""
import logging
from io import BytesIO
from xml.sax.saxutils import escape, quoteattr

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET

from . import pricing
from .product_views import NO_BRAND, _filtered

log = logging.getLogger('quotation')

PHOTO_LIMIT = 300   # photos are left out of the PDF when more products than this are exported (size and speed)


def _rows(request):
    products, q, brands, chosen, _low = _filtered(request)
    is_admin = pricing.is_price_admin(request.user)
    head = ['Model No', 'Brand', 'Description', 'Item Code', 'DIP Stock'] + (['Cost'] if is_admin else []) + ['Selling Price', 'Photo', 'Datasheet']
    info = pricing.stock_info_map()
    rows, prods = [], []
    for p in products.order_by('model_key'):
        prods.append(p)
        code, dip = pricing.stock_info(p.model_no, info)
        p.stock_item_code, p.dip_stock = code, dip
        row = [p.model_no, p.brand, p.description, code, dip]
        if is_admin:
            row.append(p.cost)
        row += [p.selling_price, 'Yes' if p.photo else 'No', 'Yes' if p.datasheet else 'No']
        rows.append(row)
    brand_text = ', '.join('(No brand)' if b == NO_BRAND else b for b in chosen) or 'All brands'
    return head, rows, brand_text, q, is_admin, prods


def _datasheet_url(request, p):
    """Login-only link to the library datasheet (staff who are logged in can open it)."""
    return request.build_absolute_uri(reverse('quotation:product_datasheet', args=[p.pk]))


def _thumb(photo):
    """A small JPEG of a product photo for the PDF, as (BytesIO, width, height), or None when it cannot be read."""
    from PIL import Image, ImageOps
    try:
        with photo.open('rb') as fh:
            img = ImageOps.exif_transpose(Image.open(fh))
            if img.mode in ('RGBA', 'LA', 'P'):
                img = img.convert('RGBA')
                bg = Image.new('RGB', img.size, 'white')
                bg.paste(img, mask=img.split()[-1])
                img = bg
            img = img.convert('RGB')
            img.thumbnail((160, 160))
            buf = BytesIO()
            img.save(buf, 'JPEG', quality=70)
    except Exception:
        log.warning('Product photo could not be read for the PDF export', exc_info=True)
        return None
    buf.seek(0)
    return buf, img.size[0], img.size[1]


def _stamp():
    return timezone.localtime().strftime('%Y%m%d_%H%M')


@login_required
@require_GET
def product_export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    head, rows, brand_text, q, is_admin, prods = _rows(request)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3A8A')
    money_cols = {i for i, h in enumerate(head, start=1) if h in ('Cost', 'Selling Price')}
    for r, prod in zip(rows, prods):
        ws.append([None] * len(r))
        n = ws.max_row
        for i, v in enumerate(r, start=1):
            cell = ws.cell(row=n, column=i)
            if v is None:
                continue
            if head[i - 1] == 'DIP Stock':
                cell.value = float(v)
                cell.number_format = '#,##0.##'
            elif i in money_cols:
                cell.value = float(v)
                cell.number_format = '#,##0.00'
            else:
                cell.value = str(v)
                cell.data_type = 's'   # always text: a value starting with = + - @ must never run as a formula
            if head[i - 1] == 'Description':
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if head[i - 1] == 'Datasheet' and prod.datasheet:
                cell.hyperlink = _datasheet_url(request, prod)   # click to open (staff must be logged in)
                cell.font = Font(color='1D4ED8', underline='single')
    widths = {'Model No': 22, 'Brand': 18, 'Description': 60, 'Item Code': 14, 'DIP Stock': 11, 'Cost': 12, 'Selling Price': 14, 'Photo': 9, 'Datasheet': 11}
    for i, h in enumerate(head, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths[h]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet('Filter')
    info.append(['Exported', timezone.localtime().strftime('%d %b %Y %H:%M')])
    info.append(['Brands', brand_text])
    info.append(['Search', q or '-'])
    info.append(['Products', len(rows)])
    info.column_dimensions['A'].width = 14
    info.column_dimensions['B'].width = 50
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="products_{_stamp()}.xlsx"'
    resp['Cache-Control'] = 'private, no-store'
    return resp


@login_required
@require_GET
def product_export_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    head, rows, brand_text, q, is_admin, prods = _rows(request)
    base = getSampleStyleSheet()['Normal']
    cell = ParagraphStyle('pc', parent=base, fontSize=8, leading=10, wordWrap='CJK')
    cell_r = ParagraphStyle('pcr', parent=cell, alignment=2)
    hcell = ParagraphStyle('ph', parent=cell, fontName='Helvetica-Bold', textColor=colors.white)
    title = ParagraphStyle('pt', parent=base, fontName='Helvetica-Bold', fontSize=14, leading=18)

    def money(v):
        return '' if v is None else f'{v:,.2f}'

    with_photos = len(prods) <= PHOTO_LIMIT
    pdf_head = ['Model No'] + (['Photo'] if with_photos else []) + ['Brand', 'Description', 'Item Code', 'DIP Stock'] + (['Cost'] if is_admin else []) + ['Selling Price', 'Datasheet']
    dash = Paragraph('-', cell)
    data = [[Paragraph(h, hcell) for h in pdf_head]]
    for p in prods:
        line = [Paragraph(escape(p.model_no), cell)]
        if with_photos:
            thumb = _thumb(p.photo) if p.photo else None
            if thumb:
                buf_img, w, h = thumb
                scale = min(16 * mm / w, 16 * mm / h)
                line.append(Image(buf_img, width=w * scale, height=h * scale))
            else:
                line.append(dash)
        line += [Paragraph(escape(p.brand or ''), cell), Paragraph(escape(p.description or '').replace('\n', '<br/>'), cell)]
        line += [Paragraph(escape(p.stock_item_code or '') or '-', cell), Paragraph(pricing.format_stock(p.dip_stock) or '-', cell_r)]
        if is_admin:
            line.append(Paragraph(money(p.cost), cell_r))
        line.append(Paragraph(money(p.selling_price), cell_r))
        if p.datasheet:
            line.append(Paragraph(f'<link href={quoteattr(_datasheet_url(request, p))} color="#1d4ed8"><u>View</u></link>', cell))
        else:
            line.append(dash)
        data.append(line)
    if not prods:
        data.append([Paragraph('No products match.', cell)] + [''] * (len(pdf_head) - 1))

    page = landscape(A4)
    usable = page[0] - 20 * mm
    fixed = {'Model No': 36, 'Photo': 20, 'Brand': 26, 'Item Code': 22, 'DIP Stock': 18, 'Cost': 20, 'Selling Price': 26, 'Datasheet': 20}
    fixed_total = sum(fixed[h] for h in pdf_head if h in fixed) * mm
    widths = [fixed[h] * mm if h in fixed else usable - fixed_total for h in pdf_head]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F3A8A')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f'Page {doc.page}')
        canvas.restoreState()

    sub = f'Brands: {brand_text}' + (f'   |   Search: {q}' if q else '') + f'   |   {len(prods)} products   |   {timezone.localtime().strftime("%d %b %Y %H:%M")}'
    note = [] if with_photos else [Paragraph(f'Photos are left out because more than {PHOTO_LIMIT} products are selected. Filter by brand to include them.', cell), Spacer(1, 2 * mm)]
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
                            title='Product Library', author='Junaid Galleria')
    doc.build([Paragraph('Product Library', title), Paragraph(escape(sub), cell), Spacer(1, 2 * mm)] + note + [Spacer(1, 2 * mm), table],
              onFirstPage=footer, onLaterPages=footer)
    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="products_{_stamp()}.pdf"'
    resp['Cache-Control'] = 'private, no-store'
    return resp
