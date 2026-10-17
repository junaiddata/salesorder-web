"""
Excel (.xlsx) export of a Bangladesh PI: the same content as the PDF, as an
editable spreadsheet. Amounts and totals are live formulas, so changing a
quantity or price in Excel updates the totals.

Read-only with respect to the PI data; nothing here is used by the PDF code
(it only re-uses the PDF's fixed wording constants).
"""
import os
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image as XLImage

from .pdf_builder import (
    BANK_LINES, CURRENCY, EXPORTER_LINES, IMPORTANT_NOTE, IMPORTER_LINES, INCOTERM,
)

THIN = Side(style='thin', color='000000')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(name='Arial', size=9, bold=True)
NORMAL = Font(name='Arial', size=9)
GREY = PatternFill('solid', fgColor='F2F2F2')
WRAP = Alignment(wrap_text=True, vertical='center')
MONEY = '#,##0.00'
QTY = '#,##0.##'

LOGO_PATH = os.path.join(os.path.dirname(__file__), 'assets', 'junaid_logo.png')   # from BD_letter.pdf
LOGO_ROWS = 3        # rows reserved for the logo at the top

# A: codes, B: description, C: qty, D: unit price, E: amount
WIDTHS = {'A': 22, 'B': 58, 'C': 13, 'D': 17, 'E': 17}


def _plain(html_text):
    return re.sub(r'<[^>]+>', '', html_text)


def _box(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.border = BOX


def build_pi_xlsx(pi, variant='customer'):
    """Return the workbook as bytes. variant: 'customer' or 'bank'."""
    bank = variant == 'bank'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Bank copy' if bank else 'Customer copy'
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    # Junaid logo at the top right (same place as on the letterhead)
    for row in range(1, LOGO_ROWS + 1):
        ws.row_dimensions[row].height = 22
    if os.path.exists(LOGO_PATH):
        logo = XLImage(LOGO_PATH)
        logo.width, logo.height = 330, 66
        ws.add_image(logo, 'C1')

    r = LOGO_ROWS + 1
    ws.merge_cells(f'A{r}:E{r}')
    ws[f'A{r}'] = 'PROFORMA INVOICE'
    ws[f'A{r}'].font = Font(name='Arial', size=12, bold=True, underline='single')
    ws[f'A{r}'].alignment = Alignment(horizontal='center')
    r += 2

    # Reference block: exporter (left) / importer (right)
    top = r
    left = [f'NO.   {pi.ref_no}', f'DATE   {pi.date.strftime("%d.%m.%Y")}', ''] + EXPORTER_LINES
    right = ['IMPORTER/CLIENT :'] + IMPORTER_LINES
    for i in range(max(len(left), len(right))):
        ws.merge_cells(f'A{r}:B{r}')
        ws.merge_cells(f'C{r}:E{r}')
        a, c = ws[f'A{r}'], ws[f'C{r}']
        a.value = left[i] if i < len(left) else None
        c.value = right[i] if i < len(right) else None
        a.font = BOLD if i < 2 or i == 3 else NORMAL
        c.font = BOLD if i == 0 or i == 1 else NORMAL
        r += 1
    _box(ws, f'A{top}:E{r - 1}')
    r += 1

    # Shipment details
    insurance = (pi.bank_insurance or pi.insurance) if bank else pi.insurance
    payment = (pi.bank_payment or pi.payment) if bank else pi.payment
    pairs = [
        ('PORT OF ORIGIN', pi.port_of_origin, 'COUNTRY OF ORIGIN', pi.country_of_origin),
        ('TRANSPORT', pi.transport, 'INSURANCE', insurance),
        ('DESTINATION', pi.destination, 'FINAL DESTINATION', pi.final_destination),
        ('SHIPMENT', pi.shipment, 'VALIDITY', pi.validity),
        ('CURRENCY', CURRENCY, 'INCOTERMS', INCOTERM),
    ]
    top = r
    for l1, v1, l2, v2 in pairs:
        ws[f'A{r}'], ws[f'B{r}'], ws[f'C{r}'], ws[f'E{r}'] = l1, v1, l2, v2
        ws.merge_cells(f'C{r}:D{r}')
        ws[f'A{r}'].font = ws[f'C{r}'].font = BOLD
        ws[f'B{r}'].font = ws[f'E{r}'].font = NORMAL
        r += 1
    ws[f'A{r}'], ws[f'B{r}'] = 'PAYMENT', payment
    ws.merge_cells(f'B{r}:E{r}')
    ws[f'A{r}'].font, ws[f'B{r}'].font = BOLD, NORMAL
    r += 1
    for rr in range(top, r):
        for col in 'ABCDE':
            ws[f'{col}{rr}'].alignment = WRAP
    _box(ws, f'A{top}:E{r - 1}')
    r += 1

    # Items table
    head = r
    for col, text in zip('ABCDE', ['CODES', 'DESCRIPTION OF GOODS', 'QUANTITY NOS.',
                                   'UNIT PRICE USD', 'TOTAL AMT. USD']):
        c = ws[f'{col}{r}']
        c.value, c.font, c.fill = text, BOLD, GREY
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r += 1
    first_item = r
    for it in pi.items.all():
        if it.kind == 'heading':
            ws.merge_cells(f'A{r}:E{r}')
            ws[f'A{r}'] = it.description or it.code
            ws[f'A{r}'].font = BOLD
        else:
            price = it.bank_unit_price if bank else it.unit_price
            ws[f'A{r}'] = it.code
            ws[f'B{r}'] = it.description
            ws[f'C{r}'] = float(it.qty)
            if it.is_foc:
                # text, so the SUM() totals simply skip it
                ws[f'D{r}'] = 'FOC'
                ws[f'E{r}'] = 'FOC'
                ws[f'D{r}'].alignment = ws[f'E{r}'].alignment = Alignment(horizontal='right', vertical='center')
            else:
                ws[f'D{r}'] = float(price or 0)
                ws[f'E{r}'] = f'=C{r}*D{r}'
            for col in 'ABCDE':
                ws[f'{col}{r}'].font = NORMAL
            ws[f'C{r}'].number_format = QTY
            ws[f'D{r}'].number_format = ws[f'E{r}'].number_format = MONEY
            if it.is_foc:
                ws[f'D{r}'].font = ws[f'E{r}'].font = BOLD
            ws[f'C{r}'].alignment = Alignment(horizontal='center', vertical='center')
        for col in 'AB':
            ws[f'{col}{r}'].alignment = WRAP
        r += 1
    last_item = r - 1

    def total_row(label, formula):
        nonlocal r
        ws.merge_cells(f'A{r}:D{r}')
        ws[f'A{r}'] = label
        ws[f'A{r}'].font = ws[f'E{r}'].font = BOLD
        ws[f'A{r}'].alignment = Alignment(horizontal='right')
        ws[f'E{r}'] = formula
        ws[f'E{r}'].number_format = MONEY
        row = r
        r += 1
        return row

    sum_formula = f'=SUM(E{first_item}:E{last_item})'
    if bank:
        grand = total_row('TOTAL CFR CHATTOGRAM , BANGLADESH.', sum_formula)
    else:
        country = (pi.country_of_origin or '').upper()
        exw = total_row(f'TOTAL EX-WORKS, {country}' if country else 'TOTAL EX-WORKS', sum_formula)
        frt = total_row('INLAND AND SEA FREIGHT CHARGES', float(pi.freight_charges or 0))
        ws[f'E{frt}'].font = BOLD
        grand = total_row('TOTAL CFR CHATTOGRAM , BANGLADESH.', f'=E{exw}+E{frt}')
    for col in 'ABCDE':
        ws[f'{col}{grand}'].fill = GREY
    _box(ws, f'A{head}:E{r - 1}')
    r += 1

    # Notes
    ws.merge_cells(f'A{r}:E{r}')
    ws[f'A{r}'] = _plain(IMPORTANT_NOTE)
    ws[f'A{r}'].font = Font(name='Arial', size=8)
    ws[f'A{r}'].alignment = WRAP
    ws.row_dimensions[r].height = 26
    r += 1
    if not bank and pi.freight_valid_till:
        ws.merge_cells(f'A{r}:E{r}')
        ws[f'A{r}'] = f'Freight Rate Valid Till {pi.freight_valid_till} only'
        ws[f'A{r}'].font = Font(name='Arial', size=8, bold=True)
        r += 1
    r += 1

    if bank and (pi.bank_terms_text or '').strip():
        ws[f'A{r}'] = 'TERMS & CONDITIONS FOR ISSUING L/C:'
        ws[f'A{r}'].font = Font(name='Arial', size=9, bold=True, underline='single')
        r += 1
        n = 0
        for line in pi.bank_terms_text.splitlines():
            line = re.sub(r'^\s*\d+\s*[.)\-:]\s*', '', line).strip()
            if line:
                n += 1
                ws.merge_cells(f'A{r}:E{r}')
                ws[f'A{r}'] = f'{n}.  {line}'
                ws[f'A{r}'].font = NORMAL
                ws[f'A{r}'].alignment = WRAP
                r += 1
        r += 1

    ws[f'A{r}'] = 'CONFIRMED AND ACCEPTED BY EXPORTER:'
    ws[f'C{r}'] = 'CONFIRMED AND ACCEPTED BY BUYER:'
    ws[f'A{r}'].font = ws[f'C{r}'].font = BOLD
    ws.row_dimensions[r + 1].height = 40
    r += 3

    ws.merge_cells(f'A{r}:E{r}')
    ws[f'A{r}'] = 'BANK ACCOUNT DETAILS [PAYMENT IN USD]'
    ws[f'A{r}'].font = BOLD
    bank_top = r
    r += 1
    for line in BANK_LINES:
        ws.merge_cells(f'A{r}:E{r}')
        ws[f'A{r}'] = line
        ws[f'A{r}'].font = NORMAL
        r += 1
    _box(ws, f'A{bank_top}:E{r - 1}')

    # Print setup: A4 portrait, fit to one page wide
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f'A1:E{r}'
    ws.sheet_view.showGridLines = False

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
