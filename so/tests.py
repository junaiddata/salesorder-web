import os
import shutil
import tempfile
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from PIL import Image

from unittest import mock

from django.core.management import call_command

from .management.commands.import_items2 import Command as ImportItemsCommand
from .models import ItemMedia, Items, ItemSyncStatus, Role

TMP_ROOT = tempfile.mkdtemp(prefix='item_catalog_test_')


def big_photo(name='p.jpg', size=(1600, 1200)):
    """A noisy (hard to compress) photo, like a real camera picture."""
    buf = BytesIO()
    Image.effect_noise(size, 80).convert('RGB').save(buf, 'JPEG', quality=95)
    return SimpleUploadedFile(name, buf.getvalue(), content_type='image/jpeg')


def small_png(name='s.png'):
    buf = BytesIO()
    Image.new('RGB', (40, 30), 'red').save(buf, 'PNG')
    return SimpleUploadedFile(name, buf.getvalue(), content_type='image/png')


@override_settings(ITEM_MEDIA_PRIVATE_ROOT=TMP_ROOT, ALABAMA_CATALOG_BRANDS=['ARISTON'],
                   ALLOWED_HOSTS=['testserver'])
class ItemCatalogTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP_ROOT, ignore_errors=True)

    def setUp(self):
        self.ariston = Items.objects.create(item_code='A1', item_description='Ariston heater', item_upvc='AR-1',
                                            item_firm='ARISTON', item_cost=10, item_price=15,
                                            dip_warehouse_stock=Decimal('5'), total_available_stock=Decimal('9'))
        self.grohe = Items.objects.create(item_code='G1', item_description='Grohe mixer', item_upvc='GR-1',
                                          item_firm='GROHE', dip_warehouse_stock=Decimal('0'))
        self.admin = self._user('boss', 'Admin', 'Junaid')
        self.sales = self._user('sam', 'Salesman', 'Junaid')
        self.alabama_admin = self._user('alabama_boss', 'Admin', 'Alabama')

    def _user(self, name, role, company):
        u = User.objects.create_user(name, password='x')
        Role.objects.create(user=u, role=role, company=company)
        return u

    def _login(self, user):
        self.client.force_login(user)

    # --- who sees what -----------------------------------------------------------------------------------------
    def test_requires_login(self):
        resp = self.client.get(reverse('item_catalog'))
        self.assertEqual(resp.status_code, 302)

    def test_junaid_sees_all_brands_and_items(self):
        self._login(self.sales)
        resp = self.client.get(reverse('item_catalog'))
        self.assertContains(resp, 'Ariston heater')
        self.assertContains(resp, 'Grohe mixer')

    def test_alabama_sees_only_allowlisted_brands(self):
        self._login(self.alabama_admin)
        resp = self.client.get(reverse('item_catalog'))
        self.assertContains(resp, 'Ariston heater')
        self.assertNotContains(resp, 'Grohe mixer')
        # a brand picked by editing the URL is still filtered out
        resp = self.client.get(reverse('item_catalog'), {'brand': 'GROHE'})
        self.assertNotContains(resp, 'Grohe mixer')

    def test_alabama_cannot_fetch_other_brand_files_or_upload(self):
        ItemMedia.objects.create(item=self.grohe, photo=None)
        self._login(self.alabama_admin)
        self.assertEqual(self.client.get(reverse('item_catalog_photo', args=[self.grohe.pk])).status_code, 404)
        resp = self.client.post(reverse('item_catalog_upload', args=[self.grohe.pk]), {'photo': small_png()})
        self.assertEqual(resp.status_code, 404)

    @override_settings(ALABAMA_CATALOG_BRANDS=[])
    def test_empty_allowlist_shows_everything_to_alabama(self):
        self._login(self.alabama_admin)
        self.assertContains(self.client.get(reverse('item_catalog')), 'Grohe mixer')

    def test_cost_is_admin_only(self):
        self._login(self.sales)
        self.assertNotContains(self.client.get(reverse('item_catalog')), '<th>Cost</th>')
        self._login(self.admin)
        self.assertContains(self.client.get(reverse('item_catalog')), '<th>Cost</th>')

    def test_hide_zero_stock_and_search(self):
        self._login(self.sales)
        resp = self.client.get(reverse('item_catalog'), {'instock': '1'})
        self.assertContains(resp, 'Ariston heater')
        self.assertNotContains(resp, 'Grohe mixer')
        resp = self.client.get(reverse('item_catalog'), {'q': 'GR-1'})
        self.assertContains(resp, 'Grohe mixer')
        self.assertNotContains(resp, 'Ariston heater')

    def test_items_without_media_still_listed(self):
        self._login(self.admin)
        self.assertEqual(self.client.get(reverse('item_catalog')).context['shown'], 2)
        self.assertEqual(self.client.get(reverse('item_catalog'), {'missing': 'photo'}).context['shown'], 2)

    # --- uploads -----------------------------------------------------------------------------------------------
    def test_non_admin_cannot_upload(self):
        self._login(self.sales)
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': small_png()})
        self.assertFalse(ItemMedia.objects.exists())

    def test_photo_is_shrunk_and_replacing_deletes_the_old_file(self):
        self._login(self.admin)
        raw = big_photo()
        raw_size = raw.size
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': raw})
        media = ItemMedia.objects.get(item=self.ariston)
        first = media.photo.path
        self.assertTrue(os.path.exists(first))
        with Image.open(first) as img:
            self.assertLessEqual(max(img.size), 800)
        self.assertLess(os.path.getsize(first), raw_size / 3)

        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': small_png()})
        media.refresh_from_db()
        self.assertNotEqual(media.photo.path, first)
        self.assertFalse(os.path.exists(first))          # old file gone, no pile-up
        self.assertEqual(ItemMedia.objects.count(), 1)   # still one row per item

    def test_invalid_upload_changes_nothing(self):
        self._login(self.admin)
        bad = SimpleUploadedFile('x.jpg', b'not an image', content_type='image/jpeg')
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': bad})
        self.assertFalse(ItemMedia.objects.exists())

    def test_oversize_photo_rejected(self):
        self._login(self.admin)
        huge = SimpleUploadedFile('h.jpg', b'0' * (5 * 1024 * 1024 + 1), content_type='image/jpeg')
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': huge})
        self.assertFalse(ItemMedia.objects.exists())

    def test_datasheet_image_becomes_jpeg_and_is_served(self):
        self._login(self.admin)
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'datasheet': small_png('d.png')})
        resp = self.client.get(reverse('item_catalog_datasheet', args=[self.ariston.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'image/jpeg')
        self.assertEqual(b''.join(resp.streaming_content)[:3], b'\xff\xd8\xff')

    def test_remove_deletes_files_and_empty_row(self):
        self._login(self.admin)
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]),
                         {'photo': small_png(), 'datasheet': small_png('d.png')})
        media = ItemMedia.objects.get(item=self.ariston)
        photo_path, sheet_path = media.photo.path, media.datasheet.path
        self.client.post(reverse('item_catalog_remove', args=[self.ariston.pk]), {'kind': 'photo'})
        self.assertFalse(os.path.exists(photo_path))
        self.assertTrue(os.path.exists(sheet_path))
        self.client.post(reverse('item_catalog_remove', args=[self.ariston.pk]), {'kind': 'datasheet'})
        self.assertFalse(os.path.exists(sheet_path))
        self.assertFalse(ItemMedia.objects.exists())

    def test_stock_sync_style_update_keeps_media(self):
        """The stock sync does bulk_update on Items only; files are in ItemMedia and must survive it."""
        self._login(self.admin)
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': small_png()})
        item = Items.objects.get(pk=self.ariston.pk)
        item.item_description = 'Ariston heater v2'
        item.item_firm = 'ARISTON'
        Items.objects.bulk_update([item], ['item_description', 'item_firm'])
        self.assertTrue(ItemMedia.objects.get(item_id=self.ariston.pk).photo)
        resp = self.client.get(reverse('item_catalog_photo', args=[self.ariston.pk]))
        self.assertEqual(resp.status_code, 200)

    def test_photo_requires_login_and_has_no_public_url(self):
        self._login(self.admin)
        self.client.post(reverse('item_catalog_upload', args=[self.ariston.pk]), {'photo': small_png()})
        self.client.logout()
        resp = self.client.get(reverse('item_catalog_photo', args=[self.ariston.pk]))
        self.assertEqual(resp.status_code, 302)
        with self.assertRaises(ValueError):
            ItemMedia.objects.get(item=self.ariston).photo.url


def api_row(code, desc='x', brand='ARISTON', dip=5):
    return {'item_code': code, 'description': desc, 'upc_code': code + '-U', 'cost_price': 1, 'manufacturer': brand,
            'minimum_selling_price': 2, 'dip_stock': dip, 'total_stock': dip}


def run_sync(rows):
    with mock.patch.object(ImportItemsCommand, '_fetch_stock_json', return_value=(rows, 'ok')):
        call_command('import_items2', verbosity=0)


@override_settings(ITEM_MEDIA_PRIVATE_ROOT=TMP_ROOT, ALLOWED_HOSTS=['testserver'], ALABAMA_CATALOG_BRANDS=[])
class ItemSyncStatusTests(TestCase):
    """Items the stock app no longer lists are marked Inactive, never deleted."""

    def setUp(self):
        self.admin = User.objects.create_user('boss', password='x')
        Role.objects.create(user=self.admin, role='Admin', company='Junaid')
        self.client.force_login(self.admin)

    def codes(self, **params):
        resp = self.client.get(reverse('item_catalog'), params)
        return {it.item_code: it for it in resp.context['page']}, resp

    def test_no_status_before_first_sync(self):
        Items.objects.create(item_code='OLD', item_description='old', item_firm='B')
        items, resp = self.client.get(reverse('item_catalog')).context['page'], None
        resp = self.client.get(reverse('item_catalog'), {'status': 'inactive'})   # ignored: nothing to compare with
        self.assertFalse(resp.context['has_status'])
        self.assertEqual(resp.context['shown'], 1)
        self.assertNotContains(resp, '<th>Status</th>')

    def test_sync_marks_listed_items_and_missing_ones_become_inactive(self):
        Items.objects.create(item_code='STALE', item_description='stale', item_firm='B', item_stock=9)   # not in the API
        run_sync([api_row('A1'), api_row('A2')])
        self.assertEqual(ItemSyncStatus.objects.count(), 2)
        items, resp = self.codes()
        self.assertTrue(resp.context['has_status'])
        # the catalog lists Active items only, and no URL option brings the others back
        self.assertEqual(set(items), {'A1', 'A2'})
        self.assertTrue(items['A1'].is_active and items['A2'].is_active)
        for param in ('all', 'inactive', 'active'):
            self.assertEqual(set(self.codes(status=param)[0]), {'A1', 'A2'}, param)
        self.assertTrue(Items.objects.filter(item_code='STALE').exists())      # kept, just not shown
        self.assertEqual((resp.context['shown'], resp.context['total']), (2, 2))   # the heading count ignores inactive items
        self.assertNotIn('B', [b for b, _ in resp.context['brand_options']])       # brand B has only an inactive item

    def test_item_dropped_from_a_later_run_turns_inactive_and_comes_back(self):
        many = [api_row(f'A{i}') for i in range(10)]
        run_sync(many)
        run_sync(many[:9])                                         # A9 no longer listed (90% of the last run)
        items, _ = self.codes()
        self.assertTrue(items['A0'].is_active)
        self.assertNotIn('A9', items)                              # no longer listed: not shown
        self.assertTrue(Items.objects.filter(item_code='A9').exists())   # never deleted
        self.assertTrue(ItemSyncStatus.objects.filter(item__item_code='A9').exists())   # remembers when it was last listed
        run_sync(many)                                             # listed again
        self.assertTrue(self.codes()[0]['A9'].is_active)
        self.assertEqual(ItemSyncStatus.objects.count(), 10)       # still one row per item

    def test_a_much_smaller_run_is_not_recorded(self):
        many = [api_row(f'A{i}') for i in range(10)]
        run_sync(many)
        run_sync(many[:2])                                         # broken / partial API reply: 20% of the last run
        self.assertEqual(len(self.codes()[0]), 10)                 # nothing was marked inactive

    def test_ignored_codes_are_not_marked(self):
        from .models import IgnoreList
        IgnoreList.objects.create(item_code='SKIP')
        run_sync([api_row('A1'), api_row('SKIP')])
        self.assertFalse(ItemSyncStatus.objects.filter(item__item_code='SKIP').exists())

    def test_sync_keeps_uploaded_files(self):
        run_sync([api_row('A1')])
        item = Items.objects.get(item_code='A1')
        self.client.post(reverse('item_catalog_upload', args=[item.pk]), {'photo': small_png()})
        run_sync([api_row('A1', desc='changed')])
        self.assertTrue(ItemMedia.objects.get(item=item).photo)

    def test_etag_is_ignored_until_the_first_marks_exist(self):
        cmd = ImportItemsCommand()
        with mock.patch.object(ImportItemsCommand, '_etag_path') as path:
            path.return_value.is_file.return_value = True
            path.return_value.read_text.return_value = 'W/"abc"'
            self.assertIsNone(cmd._read_etag())                    # no marks yet: force a full import
            run_sync([api_row('A1')])
            self.assertEqual(cmd._read_etag(), 'W/"abc"')          # afterwards the normal 304 shortcut works


@override_settings(ITEM_MEDIA_PRIVATE_ROOT=TMP_ROOT, ALLOWED_HOSTS=['testserver'])
class InactiveItemsHiddenFromPickersTests(TestCase):
    """New quotations / sales orders and the email agent only offer Active items; made documents keep theirs."""

    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        self.admin = User.objects.create_user('boss', password='x')
        Role.objects.create(user=self.admin, role='Admin', company='Junaid')
        self.client.force_login(self.admin)
        self.live = Items.objects.create(item_code='L1', item_description='Live valve', item_firm='LIVEBRAND',
                                         item_cost=1, item_price=5)
        self.dead = Items.objects.create(item_code='D1', item_description='Dead valve', item_firm='DEADBRAND',
                                         item_cost=1, item_price=5)

    def mark_active(self, *items):
        """Record a sync run that listed `items`; every other item is then Inactive."""
        from django.utils import timezone
        for it in items:
            ItemSyncStatus.objects.update_or_create(item=it, defaults={'last_seen_at': timezone.now()})

    def ids(self, url, params):
        data = self.client.get(url, params).json()
        rows = data['items'] if isinstance(data, dict) else data
        return {r['item_code'] for r in rows}

    # --- pickers -----------------------------------------------------------------------------------------------
    def test_nothing_hidden_before_the_first_sync(self):
        self.assertEqual(self.ids('/get_items_by_firm/', {'firms': 'All'}), {'L1', 'D1'})
        self.assertEqual(self.ids('/api/items-search/', {'q': 'valve'}), {'L1', 'D1'})

    def test_item_lists_hide_inactive_items(self):
        self.mark_active(self.live)
        self.assertEqual(self.ids('/get_items_by_firm/', {'firms': 'All'}), {'L1'})
        self.assertEqual(self.ids('/get_items_by_firm/', {'firms': 'DEADBRAND'}), set())
        self.assertEqual(self.ids('/api/items-search/', {'q': 'valve'}), {'L1'})
        self.assertEqual(self.ids('/api/items-search/', {'q': 'Dead'}), set())

    def test_cached_item_list_is_refreshed_by_a_new_sync_run(self):
        self.assertEqual(self.ids('/get_items_by_firm/', {'firms': 'All'}), {'L1', 'D1'})   # cached (no run yet)
        self.mark_active(self.live)                                                         # a run is recorded
        self.assertEqual(self.ids('/get_items_by_firm/', {'firms': 'All'}), {'L1'})

    def test_brand_dropdowns_hide_brands_with_only_inactive_items(self):
        self.mark_active(self.live)
        for name in ('create_quotation', 'create_sales_order'):
            firms = set(self.client.get(reverse(name)).context['firms'])
            self.assertEqual(firms, {'LIVEBRAND'}, name)

    # --- documents that already exist keep their inactive items -------------------------------------------------
    def _quotation(self):
        from .models import Customer, Quotation, QuotationItem, Salesman
        cust = Customer.objects.create(customer_code='C1', customer_name='Cust')
        sm = Salesman.objects.create(salesman_name='Sam')
        q = Quotation.objects.create(customer=cust, salesman=sm)
        QuotationItem.objects.create(quotation=q, item=self.dead, quantity=2, price=9, unit='pcs', line_total=18)
        return q, cust, sm

    def test_old_quotation_still_shows_and_saves_its_inactive_item(self):
        from .models import QuotationItem
        self.mark_active(self.live)
        q, cust, sm = self._quotation()
        resp = self.client.get(reverse('edit_quotation', args=[q.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Dead valve')                       # the existing line is still on the page
        self.assertIn('DEADBRAND', set(resp.context['firms']))        # ...with its brand still selectable
        self.assertNotIn('Dead valve', self.client.get('/api/items-search/', {'q': 'valve'}).content.decode())
        self.assertEqual(self.client.get(reverse('view_quotation_details', args=[q.pk])).status_code, 200)
        # saving the edit keeps the inactive line
        self.client.post(reverse('edit_quotation', args=[q.pk]), {
            'customer': cust.pk, 'salesman': sm.pk, 'item': [self.dead.pk], 'quantity': ['3'], 'price': ['9'], 'unit': ['pcs'],
        })
        line = QuotationItem.objects.get(quotation=q)
        self.assertEqual((line.item_id, line.quantity), (self.dead.pk, 3))

    def test_old_sales_order_still_shows_and_saves_its_inactive_item(self):
        from .models import Customer, OrderItem, Salesman, SalesOrder
        self.mark_active(self.live)
        cust = Customer.objects.create(customer_code='C2', customer_name='Cust2')
        sm = Salesman.objects.create(salesman_name='Sam2')
        so = SalesOrder.objects.create(customer=cust, salesman=sm)
        OrderItem.objects.create(order=so, item=self.dead, quantity=1, price=9, unit='pcs')
        resp = self.client.get(reverse('edit_sales_order', args=[so.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Dead valve')
        self.assertIn('DEADBRAND', set(resp.context['firms']))
        self.assertEqual(self.client.get(reverse('view_sales_order_details', args=[so.pk])).status_code, 200)
        self.client.post(reverse('edit_sales_order', args=[so.pk]), {
            'customer': cust.pk, 'salesman': sm.pk, 'item': [self.dead.pk], 'quantity': ['4'], 'price': ['9'], 'unit': ['pcs'],
        })
        self.assertEqual(OrderItem.objects.get(order=so).quantity, 4)

    # --- the email agent -----------------------------------------------------------------------------------------
    def test_agent_search_ignores_inactive_items(self):
        from emailagent import tools
        self.assertIn('Dead valve', tools.lookup_item_master.call({'description': 'valve'}))   # no run yet: all items
        self.mark_active(self.live)
        out = tools.lookup_item_master.call({'description': 'valve'})
        self.assertIn('Live valve', out)
        self.assertNotIn('Dead valve', out)

    def test_active_items_helper(self):
        from .item_status import active_items
        self.assertEqual(active_items(Items.objects.filter(item_code='D1')).count(), 1)   # fail open
        self.mark_active(self.live)
        self.assertEqual(active_items(Items.objects.filter(item_code='D1')).count(), 0)
        self.assertEqual(Items.objects.filter(item_code='D1').count(), 1)                 # the item itself is untouched


@override_settings(ITEM_MEDIA_PRIVATE_ROOT=TMP_ROOT, ALLOWED_HOSTS=['testserver'], ALABAMA_CATALOG_BRANDS=['ARISTON'])
class ItemCatalogEditDeleteExportTests(TestCase):
    """Edit / delete (only for items the stock app does not list, delete only when nothing uses the item) and the exports."""

    def setUp(self):
        from django.utils import timezone
        self.admin = self._user('boss', 'Admin', 'Junaid')
        self.cost_admin = self._user('manager', 'Admin', 'Junaid')     # may also see / change cost
        self.sales = self._user('sam', 'Salesman', 'Junaid')
        self.alabama = self._user('alabama_boss', 'Admin', 'Alabama')
        self.listed = Items.objects.create(item_code='L1', item_description='Listed heater', item_upvc='UL-1',
                                           item_firm='ARISTON', item_cost=10, item_price=15, dip_warehouse_stock=Decimal('5'))
        self.free = Items.objects.create(item_code='F1', item_description='Free mixer', item_upvc='UF-1',
                                         item_firm='GROHE', item_cost=3, item_price=7)
        ItemSyncStatus.objects.create(item=self.listed, last_seen_at=timezone.now())   # F1 was never listed -> Inactive
        self.client.force_login(self.admin)

    def _user(self, name, role, company):
        u = User.objects.create_user(name, password='x')
        Role.objects.create(user=u, role=role, company=company)
        return u

    def edit(self, item, **extra):
        data = {'item_description': 'Renamed mixer', 'item_upvc': 'UF-2', 'item_firm': 'grohe', 'item_price': '9.5', 'item_cost': '99'}
        data.update(extra)
        return self.client.post(reverse('item_catalog_edit', args=[item.pk]), data)

    # --- edit --------------------------------------------------------------------------------------------------
    def test_edit_saves_an_item_the_stock_app_does_not_list(self):
        self.edit(self.free)
        self.free.refresh_from_db()
        self.assertEqual((self.free.item_description, self.free.item_upvc, self.free.item_firm, self.free.item_price),
                         ('Renamed mixer', 'UF-2', 'GROHE', 9.5))
        self.assertEqual(self.free.item_code, 'F1')
        self.assertEqual(self.free.item_cost, 3)          # 'boss' may not change cost: the posted value is ignored

    def test_cost_can_be_changed_only_by_cost_users(self):
        self.client.force_login(self.cost_admin)
        self.edit(self.free)
        self.free.refresh_from_db()
        self.assertEqual(self.free.item_cost, 99)

    def test_edit_is_refused_for_stock_app_items_non_admins_and_before_the_first_sync(self):
        self.edit(self.listed)
        self.listed.refresh_from_db()
        self.assertEqual(self.listed.item_description, 'Listed heater')
        self.client.force_login(self.sales)
        self.edit(self.free)
        self.free.refresh_from_db()
        self.assertEqual(self.free.item_description, 'Free mixer')
        self.client.force_login(self.admin)
        ItemSyncStatus.objects.all().delete()              # no sync recorded: nothing is known, so nothing is editable
        self.edit(self.free)
        self.free.refresh_from_db()
        self.assertEqual(self.free.item_description, 'Free mixer')

    def test_invalid_edit_changes_nothing(self):
        self.edit(self.free, item_price='-1')
        self.edit(self.free, item_description='   ')
        self.free.refresh_from_db()
        self.assertEqual((self.free.item_description, self.free.item_price), ('Free mixer', 7))

    def test_edit_respects_the_alabama_brand_rule(self):
        self.client.force_login(self.alabama)
        self.assertEqual(self.edit(self.free).status_code, 404)    # GROHE is not on Alabama's list

    # --- delete ------------------------------------------------------------------------------------------------
    def test_delete_removes_an_unused_unlisted_item_and_its_files(self):
        self.client.post(reverse('item_catalog_upload', args=[self.free.pk]), {'photo': small_png()})
        path = ItemMedia.objects.get(item=self.free).photo.path
        self.assertTrue(os.path.exists(path))
        self.client.post(reverse('item_catalog_delete', args=[self.free.pk]))
        self.assertFalse(Items.objects.filter(pk=self.free.pk).exists())
        self.assertFalse(os.path.exists(path))

    def test_delete_is_refused_when_a_document_uses_the_item(self):
        from .models import Customer, Quotation, QuotationItem
        cust = Customer.objects.create(customer_code='C1', customer_name='Cust')
        q = Quotation.objects.create(customer=cust)
        QuotationItem.objects.create(quotation=q, item=self.free, quantity=1, price=7, unit='pcs', line_total=7)
        resp = self.client.post(reverse('item_catalog_delete', args=[self.free.pk]), follow=True)
        self.assertTrue(Items.objects.filter(pk=self.free.pk).exists())
        self.assertEqual(QuotationItem.objects.filter(quotation=q).count(), 1)     # the quotation is untouched
        self.assertContains(resp, 'cannot be deleted')

    def test_delete_is_refused_for_stock_app_items_non_admins_and_before_the_first_sync(self):
        self.client.post(reverse('item_catalog_delete', args=[self.listed.pk]))
        self.assertTrue(Items.objects.filter(pk=self.listed.pk).exists())
        self.client.force_login(self.sales)
        self.client.post(reverse('item_catalog_delete', args=[self.free.pk]))
        self.assertTrue(Items.objects.filter(pk=self.free.pk).exists())
        self.client.force_login(self.admin)
        ItemSyncStatus.objects.all().delete()
        self.client.post(reverse('item_catalog_delete', args=[self.free.pk]))
        self.assertTrue(Items.objects.filter(pk=self.free.pk).exists())

    def test_delete_needs_post(self):
        self.assertEqual(self.client.get(reverse('item_catalog_delete', args=[self.free.pk])).status_code, 405)
        self.assertTrue(Items.objects.filter(pk=self.free.pk).exists())

    def test_page_lists_active_items_only(self):
        resp = self.client.get(reverse('item_catalog'), {'status': 'all'})        # the option no longer exists
        self.assertContains(resp, 'Listed heater')
        self.assertNotContains(resp, 'Free mixer')                                # F1 is not listed by the stock app
        self.assertNotContains(resp, reverse('item_catalog_edit', args=[self.listed.pk]))    # active items cannot be edited
        self.assertNotContains(resp, reverse('item_catalog_delete', args=[self.listed.pk]))
        self.assertNotContains(resp, 'name="status"')
        self.assertNotContains(resp, 'name="missing"')

    # --- exports -----------------------------------------------------------------------------------------------
    def sheet(self, **params):
        from openpyxl import load_workbook
        resp = self.client.get(reverse('item_catalog_export_excel'), params)
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        wb = load_workbook(BytesIO(resp.content))
        rows = list(wb['Items'].iter_rows(values_only=True))
        return rows[0], rows[1:], wb

    def test_excel_follows_the_page_filters(self):
        head, rows, wb = self.sheet()                                  # Active items only
        self.assertEqual([r[0] for r in rows], ['L1'])
        self.assertEqual(wb['Filter']['B5'].value, 1)
        for param in ('all', 'inactive'):                              # inactive items cannot be exported either
            self.assertEqual([r[0] for r in self.sheet(status=param)[1]], ['L1'], param)
        self.assertEqual(self.sheet(q='mixer')[1], [])                 # F1 (inactive) never shows
        self.assertEqual([r[0] for r in self.sheet(q='heater')[1]], ['L1'])
        self.assertEqual([r[0] for r in self.sheet(brand='ARISTON')[1]], ['L1'])
        self.assertEqual(self.sheet(brand='GROHE')[1], [])

    def test_excel_cost_is_admin_only_and_can_be_left_out(self):
        head, rows, _ = self.sheet()
        self.assertIn('Cost', head)
        head, rows, _ = self.sheet(nocost='1')
        self.assertNotIn('Cost', head)
        self.client.force_login(self.sales)
        head, rows, _ = self.sheet()
        self.assertNotIn('Cost', head)

    def test_excel_text_cannot_run_as_a_formula(self):
        self.listed.item_description = '=HYPERLINK("http://x","y")'
        self.listed.save()
        head, rows, wb = self.sheet()
        ws = wb['Items']
        cell = [c for row in ws.iter_rows(min_row=2) for c in row if c.value == '=HYPERLINK("http://x","y")'][0]
        self.assertEqual(cell.data_type, 's')

    def test_exports_respect_the_alabama_brand_rule(self):
        self.client.force_login(self.alabama)
        head, rows, _ = self.sheet()
        self.assertEqual([r[0] for r in rows], ['L1'])
        self.assertEqual(self.client.get(reverse('item_catalog_export_pdf')).status_code, 200)

    def test_pdf_export(self):
        resp = self.client.get(reverse('item_catalog_export_pdf'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_pdf_with_a_photo(self):
        self.client.post(reverse('item_catalog_upload', args=[self.listed.pk]), {'photo': small_png()})
        resp = self.client.get(reverse('item_catalog_export_pdf'))
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_pdf_with_too_many_rows_is_refused_with_a_message(self):
        from unittest import mock
        with mock.patch('so.item_catalog_export.PDF_ROW_LIMIT', 0):
            resp = self.client.get(reverse('item_catalog_export_pdf'), follow=True)
        self.assertContains(resp, 'too many for one PDF')

    def test_exports_need_login(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse('item_catalog_export_excel')).status_code, 302)
        self.assertEqual(self.client.get(reverse('item_catalog_export_pdf')).status_code, 302)
