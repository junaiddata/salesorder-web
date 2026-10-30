"""Export the Item Catalog (with the search / brand / status filters that are on screen) to Excel or PDF.

Cost is included only for admins, exactly as on the page, and an admin can leave it out (?nocost=1) before sending the
file to someone. The Alabama brand rule applies because the rows come from the same filtered query as the page.
"""
import logging
from io import BytesIO
from xml.sax.saxutils import escape, quoteattr

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET

from quotation.product_export import _thumb

from .item_catalog_views import _describe, _filtered, _is_admin

log = logging.getLogger('so')

PHOTO_LIMIT = 300      # photos are left out of the PDF when more items than this are exported (size and speed)
PDF_ROW_LIMIT = 3000   # a PDF with more rows is refused: use Excel, or filter first


def _rows(request):
    f = _filtered(request)
    is_admin = _is_admin(request.user) and request.GET.get('nocost') != '1'
    head = (['Item Code', 'UPC / Model', 'Brand', 'Description', 'Total Stock', 'DIP Stock']
            + (['Cost'] if is_admin else []) + ['Min Price', 'Photo', 'Datasheet'] + (['Status'] if f['has_status'] else []))
    items = [_describe(it, f['has_status'], f['latest_run']) for it in f['items']]
    brand_text = ', '.join(f['chosen']) or 'All brands'
    status_text = 'Active items' if f['status'] == 'active' else ''
    return head, items, brand_text, status_text, f['q'], is_admin


def _datasheet_url(request, it):
    """Login-only link to the item's datasheet (staff who are logged in can open it)."""
    return request.build_absolute_uri(reverse('item_catalog_datasheet', args=[it.pk]))


def _stamp():
    return timezone.localtime().strftime('%Y%m%d_%H%M')


def _num(value):
    return None if value is None else float(value)


@login_required
@require_GET
def item_catalog_export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    head, items, brand_text, status_text, q, is_admin = _rows(request)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Items'
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3A8A')
    for it in items:
        values = {
            'Item Code': it.item_code, 'UPC / Model': it.item_upvc or '', 'Brand': it.item_firm or '',
            'Description': it.item_description or '', 'Total Stock': _num(it.total_available_stock),
            'DIP Stock': _num(it.dip_warehouse_stock), 'Cost': _num(it.item_cost), 'Min Price': _num(it.item_price),
            'Photo': 'Yes' if it.has_photo else 'No', 'Datasheet': 'Yes' if it.has_datasheet else 'No',
            'Status': 'Active' if it.is_active else 'Inactive',
        }
        ws.append([None] * len(head))
        n = ws.max_row
        for i, h in enumerate(head, start=1):
            v = values[h]
            cell = ws.cell(row=n, column=i)
            if v is None or v == '':
                continue
            if h in ('Total Stock', 'DIP Stock'):
                cell.value, cell.number_format = v, '#,##0.##'
            elif h in ('Cost', 'Min Price'):
                cell.value, cell.number_format = v, '#,##0.00'
            else:
                cell.value = str(v)
                cell.data_type = 's'   # always text: a value starting with = + - @ must never run as a formula
            if h == 'Description':
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if h == 'Datasheet' and it.has_datasheet:
                cell.hyperlink = _datasheet_url(request, it)   # click to open (staff must be logged in)
                cell.font = Font(color='1D4ED8', underline='single')
    widths = {'Item Code': 16, 'UPC / Model': 20, 'Brand': 18, 'Description': 60, 'Total Stock': 12, 'DIP Stock': 11,
              'Cost': 12, 'Min Price': 12, 'Photo': 9, 'Datasheet': 11, 'Status': 10}
    for i, h in enumerate(head, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths[h]
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet('Filter')
    info.append(['Exported', timezone.localtime().strftime('%d %b %Y %H:%M')])
    info.append(['Brands', brand_text])
    info.append(['Search', q or '-'])
    info.append(['Status', status_text or '-'])
    info.append(['Items', len(items)])
    info.column_dimensions['A'].width = 14
    info.column_dimensions['B'].width = 50
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="item_catalog_{_stamp()}.xlsx"'
    resp['Cache-Control'] = 'private, no-store'
    return resp


@login_required
@require_GET
def item_catalog_export_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    head, items, brand_text, status_text, q, is_admin = _rows(request)
    if len(items) > PDF_ROW_LIMIT:
        messages.error(request, f'{len(items)} items are too many for one PDF (max {PDF_ROW_LIMIT}). '
                                'Filter by brand first, or download the Excel file instead.')
        return redirect(reverse('item_catalog') + ('?' + request.GET.urlencode() if request.GET else ''))

    base = getSampleStyleSheet()['Normal']
    cell = ParagraphStyle('ic', parent=base, fontSize=8, leading=10, wordWrap='CJK')
    cell_r = ParagraphStyle('icr', parent=cell, alignment=2)
    hcell = ParagraphStyle('ih', parent=cell, fontName='Helvetica-Bold', textColor=colors.white)
    title = ParagraphStyle('it', parent=base, fontName='Helvetica-Bold', fontSize=14, leading=18)

    def money(v):
        return '' if v is None else f'{v:,.2f}'

    with_photos = len(items) <= PHOTO_LIMIT
    pdf_head = (['Item Code'] + (['Photo'] if with_photos else [])
                + ['UPC / Model', 'Brand', 'Description', 'Total Stock', 'DIP Stock'] + (['Cost'] if is_admin else [])
                + ['Min Price', 'Datasheet'] + (['Status'] if 'Status' in head else []))
    dash = Paragraph('-', cell)
    data = [[Paragraph(h, hcell) for h in pdf_head]]
    for it in items:
        line = [Paragraph(escape(it.item_code), cell)]
        if with_photos:
            thumb = _thumb(it.media.photo) if it.has_photo else None
            if thumb:
                buf_img, w, h = thumb
                scale = min(16 * mm / w, 16 * mm / h)
                line.append(Image(buf_img, width=w * scale, height=h * scale))
            else:
                line.append(dash)
        line += [Paragraph(escape(it.item_upvc or '') or '-', cell), Paragraph(escape(it.item_firm or ''), cell),
                 Paragraph(escape(it.item_description or '').replace('\n', '<br/>'), cell),
                 Paragraph(it.total_stock_text or '-', cell_r), Paragraph(it.dip_stock_text or '-', cell_r)]
        if is_admin:
            line.append(Paragraph(money(it.item_cost), cell_r))
        line.append(Paragraph(money(it.item_price), cell_r))
        if it.has_datasheet:
            line.append(Paragraph(f'<link href={quoteattr(_datasheet_url(request, it))} color="#1d4ed8"><u>View</u></link>', cell))
        else:
            line.append(dash)
        if 'Status' in head:
            line.append(Paragraph('Active' if it.is_active else 'Inactive', cell))
        data.append(line)
    if not items:
        data.append([Paragraph('No items match.', cell)] + [''] * (len(pdf_head) - 1))

    page = landscape(A4)
    usable = page[0] - 20 * mm
    fixed = {'Item Code': 24, 'Photo': 20, 'UPC / Model': 30, 'Brand': 24, 'Total Stock': 17, 'DIP Stock': 16, 'Cost': 18,
             'Min Price': 20, 'Datasheet': 18, 'Status': 16}
    fixed_total = sum(fixed[h] for h in pdf_head if h in fixed) * mm
    widths = [fixed[h] * mm if h in fixed else usable - fixed_total for h in pdf_head]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F3A8A')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f'Page {doc.page}')
        canvas.restoreState()

    sub = (f'Brands: {brand_text}' + (f'   |   {status_text}' if status_text else '') + (f'   |   Search: {q}' if q else '')
           + f'   |   {len(items)} items   |   {timezone.localtime().strftime("%d %b %Y %H:%M")}')
    note = [] if with_photos else [Paragraph(f'Photos are left out because more than {PHOTO_LIMIT} items are selected. Filter by brand to include them.', cell), Spacer(1, 2 * mm)]
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=12 * mm,
                            title='Item Catalog', author='Junaid')
    doc.build([Paragraph('Item Catalog', title), Paragraph(escape(sub), cell), Spacer(1, 2 * mm)] + note + [Spacer(1, 2 * mm), table],
              onFirstPage=footer, onLaterPages=footer)
    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="item_catalog_{_stamp()}.pdf"'
    resp['Cache-Control'] = 'private, no-store'
    return resp
