"""
Supplier / Brandwise Item Analysis
==================================
Compares the brands (Items.item_firm) the user selects SIDE BY SIDE, using the
same data and net-sales definition as Item Analysis (AR Invoice lines + AR
Credit Memo lines, credit memos netting negative):

  - Summary comparison : one column per brand (Sales, GP, GP %, Quantity,
                         Items, Avg Rate, Share of selected brands)
  - Item comparison    : one block of columns per brand, each listing that
                         brand's top items by sales, ranked side by side

Filters: Year, Store, Month (multi), Salesman (multi), Brand (multi), Item search.
All items of the selected brands are listed, paginated.
"""
from collections import defaultdict
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db.models import Sum, Max, Q, DecimalField, Value
from django.db.models.functions import Coalesce
from django.core.paginator import Paginator
from django.shortcuts import render

from .brandwise_sales_analysis_views import (
    MONTH_NAMES_SHORT,
    _net_value_expr,
    _gp_value_expr,
    _pct,
    _user_is_admin,
)
from .models import (
    Items,
    SAPARInvoice,
    SAPARInvoiceItem,
    SAPARCreditMemo,
    SAPARCreditMemoItem,
)
from .sap_salesorder_views import salesman_scope_q_salesorder

PAGE_SIZE = 50  # ranked rows per page (each row holds one item per brand)


def _qty_expr():
    return Coalesce(Sum('quantity'), Value(0, output_field=DecimalField()))


def _analysis_data(request):
    """Apply the page filters and build the summary + full ranked rows (unpaginated)."""
    today = date.today()
    current_year = today.year

    # ── Filters ──────────────────────────────────────────────────
    calendar_years = list(range(2024, current_year + 2))
    try:
        selected_year = int(request.GET.get('year', '').strip() or current_year)
    except (ValueError, TypeError):
        selected_year = current_year
    if selected_year not in calendar_years:
        selected_year = current_year

    store_filter = (request.GET.get('store', 'Total') or 'Total').strip()
    if store_filter not in ('HO', 'Others', 'Total'):
        store_filter = 'Total'

    def _parse_month(val, default):
        try:
            m = int(str(val).strip())
        except (ValueError, TypeError):
            return default
        return m if 1 <= m <= 12 else default

    # Multi-select months; empty = all months of the year
    selected_months = sorted({
        m for m in (_parse_month(v, None) for v in request.GET.getlist('month')) if m
    })

    selected_salesmen = [s.strip() for s in request.GET.getlist('salesman') if s.strip()]
    selected_brands = list(dict.fromkeys(
        b.strip() for b in request.GET.getlist('brand') if b and b.strip()
    ))
    search_query = request.GET.get('q', '').strip()
    hide_zero = request.GET.get('hide_zero') == '1'  # Exclude items whose net sales is 0

    is_admin = _user_is_admin(request.user)

    # ── Option lists ─────────────────────────────────────────────
    scope_q = salesman_scope_q_salesorder(request.user)
    salesmen = list(
        SAPARInvoice.objects.filter(scope_q)
        .exclude(salesman_name__isnull=True)
        .exclude(salesman_name='')
        .exclude(salesman_name__iexact='Z.DUTY')
        .values_list('salesman_name', flat=True)
        .distinct()
        .order_by('salesman_name')
    )
    brands = list(
        Items.objects.exclude(item_firm__isnull=True)
        .exclude(item_firm='')
        .values_list('item_firm', flat=True)
        .distinct()
        .order_by('item_firm')
    )

    brand_summaries = []
    rank_rows = []
    grand_sales = Decimal('0')

    if selected_brands:
        inv_headers = (
            SAPARInvoice.objects.filter(scope_q)
            .exclude(salesman_name__iexact='Z.DUTY')
            .filter(posting_date__year=selected_year)
        )
        cm_headers = (
            SAPARCreditMemo.objects.filter(scope_q)
            .exclude(salesman_name__iexact='Z.DUTY')
            .filter(posting_date__year=selected_year)
        )
        if selected_months:
            inv_headers = inv_headers.filter(posting_date__month__in=selected_months)
            cm_headers = cm_headers.filter(posting_date__month__in=selected_months)
        if store_filter in ('HO', 'Others'):
            inv_headers = inv_headers.filter(store=store_filter)
            cm_headers = cm_headers.filter(store=store_filter)
        if selected_salesmen:
            inv_headers = inv_headers.filter(salesman_name__in=selected_salesmen)
            cm_headers = cm_headers.filter(salesman_name__in=selected_salesmen)

        inv_items = SAPARInvoiceItem.objects.filter(
            invoice__in=inv_headers, item__item_firm__in=selected_brands,
        ).exclude(item_code__isnull=True).exclude(item_code='')
        cm_items = SAPARCreditMemoItem.objects.filter(
            credit_memo__in=cm_headers, item__item_firm__in=selected_brands,
        ).exclude(item_code__isnull=True).exclude(item_code='')

        if search_query:
            text_q = (Q(item_code__icontains=search_query)
                      | Q(item_description__icontains=search_query)
                      | Q(upc_code__icontains=search_query))
            inv_items = inv_items.filter(text_q)
            cm_items = cm_items.filter(text_q)

        # brand -> item_code -> aggregates
        data = defaultdict(dict)

        def _accumulate(rows):
            for r in rows:
                brand = r['item__item_firm']
                code = r['item_code']
                d = data[brand].setdefault(code, {
                    'item_code': code,
                    'description': r['description'] or '',
                    'sales': Decimal('0'), 'gp': Decimal('0'), 'qty': Decimal('0'),
                })
                d['sales'] += r['sales'] or Decimal('0')
                d['gp'] += r['gp'] or Decimal('0')
                d['qty'] += r['qty'] or Decimal('0')
                if not d['description'] and r['description']:
                    d['description'] = r['description']

        _accumulate(
            inv_items.values('item__item_firm', 'item_code').annotate(
                sales=_net_value_expr(), gp=_gp_value_expr(), qty=_qty_expr(),
                description=Max('item_description'))
        )
        _accumulate(
            cm_items.values('item__item_firm', 'item_code').annotate(
                sales=_net_value_expr(), gp=_gp_value_expr(), qty=_qty_expr(),
                description=Max('item_description'))
        )

        # Stock on hand (Items.total_available_stock): brand total over ALL of the
        # brand's items, plus a per-item lookup for the ranked lists.
        stock_by_brand = {
            row['item_firm']: row['stock'] or Decimal('0')
            for row in Items.objects.filter(item_firm__in=selected_brands)
            .values('item_firm').annotate(stock=Sum('total_available_stock'))
        }
        ranked_codes = set()
        for brand_items in data.values():
            ranked_codes.update(brand_items.keys())
        item_stock = {
            code: stk
            for code, stk in Items.objects.filter(item_code__in=ranked_codes)
            .values_list('item_code', 'total_available_stock')
        }

        # Per-brand summary + ranked item lists (columns follow selection order)
        ranked = {}
        for brand in selected_brands:
            brand_items = data.get(brand, {}).values()
            if hide_zero:
                # Drop 0-net-sales items before ranking so the list closes up
                brand_items = [i for i in brand_items if i['sales'] != 0]
            items = sorted(brand_items, key=lambda x: x['sales'], reverse=True)
            ranked[brand] = items
            sales = sum((i['sales'] for i in items), Decimal('0'))
            gp = sum((i['gp'] for i in items), Decimal('0'))
            qty = sum((i['qty'] for i in items), Decimal('0'))
            grand_sales += sales
            brand_summaries.append({
                'brand': brand,
                'sales': sales,
                'gp': gp,
                'gp_pct': _pct(gp, sales),
                'qty': qty,
                'item_count': len(items),
                'stock': stock_by_brand.get(brand, Decimal('0')),
                'avg_rate': (sales / qty) if qty else Decimal('0'),
            })
        for s in brand_summaries:
            s['share'] = _pct(s['sales'], grand_sales)

        # Side-by-side rank table: row i holds each brand's i-th best item
        max_items = max((len(v) for v in ranked.values()), default=0)
        for i in range(max_items):
            cells = []
            for brand in selected_brands:
                items = ranked[brand]
                if i < len(items):
                    it = items[i]
                    cells.append({
                        'item_code': it['item_code'],
                        'description': it['description'],
                        'qty': it['qty'],
                        'stock': item_stock.get(it['item_code']),
                        'sales': it['sales'],
                        'gp': it['gp'],
                        'gp_pct': _pct(it['gp'], it['sales']),
                    })
                else:
                    cells.append(None)
            rank_rows.append({'rank': i + 1, 'cells': cells})

    return {
        'calendar_years': calendar_years,
        'selected_year': selected_year,
        'store_filter': store_filter,
        'selected_months': selected_months,
        'salesmen': salesmen,
        'selected_salesmen': selected_salesmen,
        'brands': brands,
        'selected_brands': selected_brands,
        'search_query': search_query,
        'hide_zero': hide_zero,
        'is_admin': is_admin,
        'brand_summaries': brand_summaries,
        'rank_rows': rank_rows,
        'grand_sales': grand_sales,
    }


@login_required
def supplier_brandwise_item_analysis(request):
    d = _analysis_data(request)
    rank_rows = d.pop('rank_rows')

    params = request.GET.copy()
    params.pop('page', None)
    paginator = Paginator(rank_rows, PAGE_SIZE)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        **d,
        'all_months': [{'num': i + 1, 'name': MONTH_NAMES_SHORT[i]} for i in range(12)],
        'page_obj': page_obj,
        'total_ranked_rows': len(rank_rows),
        'page_querystring': params.urlencode(),
        'rank_rows': page_obj.object_list if page_obj else [],
    }
    return render(request, 'salesorders/supplier_brandwise_item_analysis.html', context)


def _filter_summary(d):
    months = ', '.join(MONTH_NAMES_SHORT[m - 1] for m in d['selected_months']) or 'All months'
    salesmen = ', '.join(d['selected_salesmen']) or 'All salesmen'
    parts = [
        f"Year: {d['selected_year']}",
        f"Store: {d['store_filter']}",
        f"Months: {months}",
        f"Salesmen: {salesmen}",
        f"Brands: {', '.join(d['selected_brands']) or '-'}",
    ]
    if d['search_query']:
        parts.append(f"Search: {d['search_query']}")
    return parts


def _num(v, places=2):
    return f"{v:,.{places}f}" if v is not None else '-'


@login_required
def export_supplier_brandwise_item_analysis_excel(request):
    from io import BytesIO
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    d = _analysis_data(request)
    is_admin = d['is_admin']
    summaries = d['brand_summaries']

    wb = Workbook()
    ws = wb.active
    ws.title = 'Brand Summary'

    head_fill = PatternFill('solid', fgColor='1E3A8A')
    group_fill = PatternFill('solid', fgColor='3B82F6')
    head_font = Font(bold=True, color='FFFFFF')
    thin = Side(style='thin', color='CBD5E1')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')

    ws.cell(row=1, column=1, value='Supplier/Brandwise Item Analysis').font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=' | '.join(_filter_summary(d)))

    # Summary sheet: metric rows x brand columns
    r = 4
    ws.cell(row=r, column=1, value='Metric')
    for i, b in enumerate(summaries):
        ws.cell(row=r, column=2 + i, value=b['brand'])
    for c in range(1, 2 + len(summaries)):
        cell = ws.cell(row=r, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border

    metrics = [('Sales', 'sales', '#,##0.00')]
    if is_admin:
        metrics += [('GP', 'gp', '#,##0.00'), ('GP %', 'gp_pct', '0.00"%"')]
    metrics += [
        ('Stock', 'stock', '#,##0.##'),
        ('Quantity', 'qty', '#,##0.##'),
        ('Items', 'item_count', '#,##0'),
        ('Avg Rate', 'avg_rate', '#,##0.00'),
        ('Share %', 'share', '0.00"%"'),
    ]
    for label, key, fmt in metrics:
        r += 1
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=1).border = border
        for i, b in enumerate(summaries):
            cell = ws.cell(row=r, column=2 + i, value=float(b[key]))
            cell.number_format = fmt
            cell.border = border
    ws.column_dimensions['A'].width = 16
    for i in range(len(summaries)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 22

    # Item comparison sheet: one block of columns per brand
    ws2 = wb.create_sheet('Item Comparison')
    sub_cols = ['Item Code', 'Description', 'Qty', 'Stock', 'Sales']
    if is_admin:
        sub_cols += ['GP', 'GP %']
    n = len(sub_cols)

    if summaries:
        ws2.cell(row=1, column=1, value='Rank')
        ws2.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        for bi, b in enumerate(summaries):
            c0 = 2 + bi * n
            ws2.cell(row=1, column=c0, value=b['brand'])
            ws2.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + n - 1)
            for si, name in enumerate(sub_cols):
                ws2.cell(row=2, column=c0 + si, value=name)
        for rr in (1, 2):
            for c in range(1, 2 + len(summaries) * n):
                cell = ws2.cell(row=rr, column=c)
                cell.fill = group_fill if rr == 1 else head_fill
                cell.font, cell.alignment, cell.border = head_font, center, border

    row_i = 2
    for row in d['rank_rows']:
        row_i += 1
        ws2.cell(row=row_i, column=1, value=row['rank']).border = border
        for bi, c in enumerate(row['cells']):
            c0 = 2 + bi * n
            if c:
                vals = [c['item_code'], c['description'], float(c['qty']),
                        float(c['stock']) if c['stock'] is not None else None,
                        float(c['sales'])]
                if is_admin:
                    vals += [float(c['gp']), float(c['gp_pct'])]
            else:
                vals = [None] * n
            for si, v in enumerate(vals):
                cell = ws2.cell(row=row_i, column=c0 + si, value=v)
                cell.border = border
                name = sub_cols[si]
                if name in ('Qty', 'Stock'):
                    cell.number_format = '#,##0.##'
                elif name in ('Sales', 'GP'):
                    cell.number_format = '#,##0.00'
                elif name == 'GP %':
                    cell.number_format = '0.00"%"'
    ws2.column_dimensions['A'].width = 7
    widths = {'Item Code': 16, 'Description': 34, 'Qty': 10, 'Stock': 10,
              'Sales': 14, 'GP': 12, 'GP %': 8}
    for bi in range(len(summaries)):
        for si, name in enumerate(sub_cols):
            ws2.column_dimensions[get_column_letter(2 + bi * n + si)].width = widths[name]
    ws2.freeze_panes = 'B3'

    out = BytesIO()
    wb.save(out)
    resp = HttpResponse(
        out.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = 'attachment; filename="supplier_brandwise_item_analysis.xlsx"'
    return resp


@login_required
def export_supplier_brandwise_item_analysis_pdf(request):
    from io import BytesIO
    from django.http import HttpResponse
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    d = _analysis_data(request)
    is_admin = d['is_admin']
    summaries = d['brand_summaries']

    buf = BytesIO()
    page_w = landscape(A4)[0]
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=20, rightMargin=20,
                            topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    small = ParagraphStyle('small', parent=styles['Normal'], fontSize=7, leading=8.5)
    story = [
        Paragraph('Supplier/Brandwise Item Analysis', styles['Title']),
        Paragraph(' | '.join(_filter_summary(d)), small),
        Spacer(1, 10),
    ]

    navy = colors.HexColor('#1E3A8A')
    blue = colors.HexColor('#3B82F6')
    grid = colors.HexColor('#CBD5E1')

    # Summary: metric rows x brand columns
    rows = [['Metric'] + [b['brand'] for b in summaries]]
    rows.append(['Sales'] + [_num(b['sales']) for b in summaries])
    if is_admin:
        rows.append(['GP'] + [_num(b['gp']) for b in summaries])
        rows.append(['GP %'] + [f"{_num(b['gp_pct'])}%" for b in summaries])
    rows.append(['Stock'] + [_num(b['stock'], 0) for b in summaries])
    rows.append(['Quantity'] + [_num(b['qty'], 0) for b in summaries])
    rows.append(['Items'] + [str(b['item_count']) for b in summaries])
    rows.append(['Avg Rate'] + [_num(b['avg_rate']) for b in summaries])
    rows.append(['Share %'] + [f"{_num(b['share'])}%" for b in summaries])
    t = Table(rows, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), navy), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7), ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.4, grid),
    ]))
    story += [t, Spacer(1, 14)]

    # Item comparison: brands are split into chunks so the columns stay readable
    sub_cols = ['Item', 'Qty', 'Stock', 'Sales'] + (['GP', 'GP %'] if is_admin else [])
    n = len(sub_cols)
    per_chunk = 2 if is_admin else 3
    cell_style = ParagraphStyle('cell', parent=small, fontSize=6.5, leading=7.5)

    def _cell_vals(c):
        if not c:
            return [''] * n
        label = Paragraph(f"<b>{c['item_code']}</b><br/>{(c['description'] or '')[:40]}", cell_style)
        vals = [label, _num(c['qty'], 0),
                _num(c['stock'], 0) if c['stock'] is not None else '-', _num(c['sales'])]
        if is_admin:
            vals += [_num(c['gp']), f"{_num(c['gp_pct'])}%"]
        return vals

    if d['rank_rows']:
        for start in range(0, len(summaries), per_chunk):
            chunk = range(start, min(start + per_chunk, len(summaries)))
            top = ['#']
            sub = ['']
            for bi in chunk:
                top += [summaries[bi]['brand']] + [''] * (n - 1)
                sub += sub_cols
            data = [top, sub]
            for row in d['rank_rows']:
                line = [str(row['rank'])]
                for bi in chunk:
                    line += _cell_vals(row['cells'][bi])
                data.append(line)
            first_w = 18
            block_w = (page_w - 40 - first_w) / len(chunk)
            item_w = block_w * 0.34
            other_w = (block_w - item_w) / (n - 1)
            widths = [first_w] + ([item_w] + [other_w] * (n - 1)) * len(chunk)
            t = Table(data, colWidths=widths, repeatRows=2)
            style = [
                ('BACKGROUND', (0, 0), (-1, 0), blue), ('BACKGROUND', (0, 1), (-1, 1), navy),
                ('TEXTCOLOR', (0, 0), (-1, 1), colors.white),
                ('FONTNAME', (0, 0), (-1, 1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 6.5), ('GRID', (0, 0), (-1, -1), 0.3, grid),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ]
            for k in range(len(chunk)):
                c0 = 1 + k * n
                style.append(('SPAN', (c0, 0), (c0 + n - 1, 0)))
                style.append(('ALIGN', (c0, 0), (c0 + n - 1, 0), 'CENTER'))
                style.append(('ALIGN', (c0 + 1, 2), (c0 + n - 1, -1), 'RIGHT'))
            t.setStyle(TableStyle(style))
            story += [Paragraph('Item Comparison', styles['Heading3']), t, Spacer(1, 12)]

    doc.build(story)
    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = 'attachment; filename="supplier_brandwise_item_analysis.pdf"'
    return resp
